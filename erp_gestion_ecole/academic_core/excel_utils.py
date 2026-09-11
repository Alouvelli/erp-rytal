"""
Utilitaires partagés pour l'import en masse depuis un fichier Excel (.xlsx) :
génération d'un modèle téléchargeable stylé, lecture d'un classeur uploadé.
Voir academic_core/pdf_utils.py pour l'équivalent côté génération de PDF.
"""
import io

from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', fgColor='003D82')
HEADER_FONT = Font(color='FFFFFF', bold=True, size=11)
EXAMPLE_FONT = Font(italic=True, color='15803D')
_THIN = Side(style='thin', color='CBD5E1')
CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def build_xlsx_template(*, filename, sheet_title, headers, example_row=None, notes=None, blank_rows=200):
    """
    Génère et retourne (HttpResponse en pièce jointe) un classeur Excel stylé
    pour l'import en masse : ligne d'en-têtes (bleu marine), une ligne
    d'exemple optionnelle (vert italique) à ne pas modifier/supprimer — voir
    `read_xlsx_upload`, qui ignore systématiquement les 2 premières lignes —
    colonnes dimensionnées, volets figés sous les en-têtes. `notes`, si
    fourni, est une liste de lignes ajoutées dans un second onglet "Notes"
    (ex. valeurs autorisées pour un champ à choix).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = CELL_BORDER
        ws.column_dimensions[get_column_letter(idx)].width = max(18, len(header) + 6)
    ws.row_dimensions[1].height = 32

    next_row = 2
    if example_row:
        for idx, val in enumerate(example_row, start=1):
            cell = ws.cell(row=2, column=idx, value=val)
            cell.font = EXAMPLE_FONT
            cell.border = CELL_BORDER
        next_row = 3

    for r in range(next_row, next_row + blank_rows):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = CELL_BORDER

    ws.freeze_panes = f'A{next_row}'

    if notes:
        notes_ws = wb.create_sheet('Notes')
        notes_ws.column_dimensions['A'].width = 100
        for i, note in enumerate(notes, start=1):
            notes_ws.cell(row=i, column=1, value=note)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def read_xlsx_upload(uploaded_file, num_columns, start_row=3):
    """
    Lit un classeur .xlsx uploadé produit par `build_xlsx_template` (ligne 1 =
    en-têtes, ligne 2 = exemple, données à partir de `start_row`). Retourne
    une liste de `(numero_de_ligne, tuple_de_valeurs)` pour chaque ligne non
    entièrement vide.
    """
    wb = load_workbook(uploaded_file, data_only=True)
    ws = wb.active
    rows = []
    for i, row in enumerate(
        ws.iter_rows(min_row=start_row, max_col=num_columns, values_only=True), start=start_row
    ):
        if all(v in (None, '') for v in row):
            continue
        rows.append((i, row))
    return rows
