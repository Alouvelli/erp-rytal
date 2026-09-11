import io
from datetime import date, datetime
from collections import defaultdict

from django.http import HttpResponse
from django.shortcuts import render

from .models import CommunityServiceActivity
from .views import _sc_required


def _parse_date(raw, default=None):
    try:
        return datetime.strptime(raw, '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return default


def _resolve_report_filters(request):
    today = date.today()
    date_from = _parse_date(request.GET.get('from'), default=date(today.year, 1, 1))
    date_to = _parse_date(request.GET.get('to'), default=today)
    categorie = request.GET.get('categorie', '')
    statut = request.GET.get('statut', '')

    qs = CommunityServiceActivity.objects.select_related('responsable').filter(
        date_debut__gte=date_from, date_debut__lte=date_to,
    ).order_by('date_debut')
    if categorie:
        qs = qs.filter(categorie=categorie)
    if statut:
        qs = qs.filter(statut=statut)
    return qs, date_from, date_to, categorie, statut


def _build_report_data(request):
    qs, date_from, date_to, categorie, statut = _resolve_report_filters(request)
    activities = list(qs)

    par_categorie = defaultdict(int)
    total_participants = 0
    for a in activities:
        par_categorie[a.get_categorie_display()] += 1
        total_participants += a.nombre_participants or 0

    return {
        'activities': activities,
        'date_from': date_from,
        'date_to': date_to,
        'categorie': categorie,
        'statut': statut,
        'categorie_label': dict(CommunityServiceActivity.CATEGORIE_CHOICES).get(categorie, ''),
        'statut_label': dict(CommunityServiceActivity.STATUT_CHOICES).get(statut, ''),
        'total_activites': len(activities),
        'total_participants': total_participants,
        'par_categorie': sorted(par_categorie.items()),
    }


@_sc_required
def activity_report(request):
    data = _build_report_data(request)
    data['categorie_choices'] = CommunityServiceActivity.CATEGORIE_CHOICES
    data['statut_choices'] = CommunityServiceActivity.STATUT_CHOICES
    return render(request, 'community_service/activity_report.html', data)


def _report_rows_for_export(data):
    rows = []
    for a in data['activities']:
        rows.append([
            a.date_debut.strftime('%d/%m/%Y'),
            a.titre,
            a.get_categorie_display(),
            a.lieu or '—',
            a.beneficiaires or '—',
            str(a.nombre_participants) if a.nombre_participants is not None else '—',
            a.get_statut_display(),
            a.responsable.get_full_name() if a.responsable else '—',
        ])
    return rows


_EXPORT_HEADERS = ['Date', 'Titre', 'Catégorie', 'Lieu', 'Bénéficiaires', 'Participants', 'Statut', 'Responsable']


@_sc_required
def activity_report_export_excel(request):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    data = _build_report_data(request)
    rows = _report_rows_for_export(data)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapport d'activités"

    navy = "00173B"
    thin = Side(style='thin', color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ncols = len(_EXPORT_HEADERS)
    end_col = get_column_letter(ncols)
    ws.merge_cells(f'A1:{end_col}1')
    ws['A1'] = f"Rapport d'activités — Service à la Communauté"
    ws['A1'].font = Font(bold=True, color="FFFFFF", size=13)
    ws['A1'].fill = PatternFill("solid", fgColor=navy)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26

    ws.merge_cells(f'A2:{end_col}2')
    ws['A2'] = f"Période : {data['date_from'].strftime('%d/%m/%Y')} — {data['date_to'].strftime('%d/%m/%Y')}"
    ws['A2'].font = Font(italic=True, color="64748B", size=9)
    ws['A2'].alignment = Alignment(horizontal='center')

    for i, h in enumerate(_EXPORT_HEADERS, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor="1E3A5F")
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border
    col_widths = [12, 32, 22, 20, 26, 13, 14, 22]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    rownum = 5
    for row in rows:
        for i, val in enumerate(row, 1):
            c = ws.cell(row=rownum, column=i, value=val)
            c.border = border
            c.alignment = Alignment(horizontal='center' if i in (1, 6, 7) else 'left', vertical='center')
        rownum += 1

    c = ws.cell(row=rownum, column=1, value=f"TOTAL : {data['total_activites']} activité(s), {data['total_participants']} participant(s)")
    ws.merge_cells(start_row=rownum, start_column=1, end_row=rownum, end_column=ncols)
    c.font = Font(bold=True, color="166534")
    c.fill = PatternFill("solid", fgColor="DCFCE7")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="rapport_activites_communaute.xlsx"'
    return response


@_sc_required
def activity_report_export_pdf(request):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    data = _build_report_data(request)
    rows = _report_rows_for_export(data)

    navy = colors.HexColor('#00173B')
    green = colors.HexColor('#166534')
    green_bg = colors.HexColor('#DCFCE7')

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1 * cm, rightMargin=1 * cm,
                             topMargin=1 * cm, bottomMargin=1 * cm)
    story = [
        Paragraph("Rapport d'activités — Service à la Communauté",
                  ParagraphStyle('t', fontSize=14, textColor=navy, alignment=TA_CENTER, spaceAfter=4)),
        Paragraph(f"Période : {data['date_from'].strftime('%d/%m/%Y')} — {data['date_to'].strftime('%d/%m/%Y')}",
                  ParagraphStyle('s', fontSize=9, textColor=colors.HexColor('#64748B'), alignment=TA_CENTER, spaceAfter=10)),
    ]

    table_data = [_EXPORT_HEADERS] + rows
    t = Table(table_data, colWidths=[2.2 * cm, 5.5 * cm, 3.5 * cm, 3 * cm, 4 * cm, 2.3 * cm, 2.3 * cm, 3.5 * cm], repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), navy), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 8), ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.3 * cm))
    total_t = Table(
        [[f"TOTAL : {data['total_activites']} activité(s), {data['total_participants']} participant(s)"]],
        colWidths=[26.3 * cm],
    )
    total_t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), green_bg), ('TEXTCOLOR', (0, 0), (-1, -1), green),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('LEFTPADDING', (0, 0), (-1, -1), 8), ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(total_t)

    doc.build(story)
    buf.seek(0)
    response = HttpResponse(buf.read(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="rapport_activites_communaute.pdf"'
    return response


@_sc_required
def activity_report_export_word(request):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    data = _build_report_data(request)
    rows = _report_rows_for_export(data)

    def _shd(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), hex_color)
        tcPr.append(shd)

    def _wc(cell, text, bold=False, size=8, color='000000', align='left'):
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == 'center' else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(str(text))
        run.bold = bold
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Rapport d'activités — Service à la Communauté")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string('00173B')

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = subtitle.add_run(f"Période : {data['date_from'].strftime('%d/%m/%Y')} — {data['date_to'].strftime('%d/%m/%Y')}")
    srun.italic = True
    srun.font.size = Pt(9)
    srun.font.color.rgb = RGBColor.from_string('64748B')

    table = doc.add_table(rows=1 + len(rows), cols=len(_EXPORT_HEADERS))
    table.style = 'Table Grid'
    for i, h in enumerate(_EXPORT_HEADERS):
        _wc(table.rows[0].cells[i], h, bold=True, color='FFFFFF', align='center')
        _shd(table.rows[0].cells[i], '1E3A5F')
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            _wc(table.rows[1 + r_idx].cells[c_idx], val, align='center' if c_idx in (0, 5, 6) else 'left')

    doc.add_paragraph()
    total_p = doc.add_paragraph()
    trun = total_p.add_run(f"TOTAL : {data['total_activites']} activité(s), {data['total_participants']} participant(s)")
    trun.bold = True
    trun.font.color.rgb = RGBColor.from_string('166534')

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = 'attachment; filename="rapport_activites_communaute.docx"'
    return response
