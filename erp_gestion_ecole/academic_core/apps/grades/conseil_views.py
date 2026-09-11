"""
Tableau de Conseil — vues de délibération semestrielle par classe.
"""
import io
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from decimal import Decimal

from .models import UniteEnseignement
from academic_core.apps.academic_structure.models import Semester, Class
from academic_core.apps.subjects.models import natural_sort_key


# ── Accès ─────────────────────────────────────────────────────────────────────

def _can_access(user):
    return (
        user.is_admin()
        or user.is_responsable()
        or user.can_manage_dept()
        or user.is_charge_examens_concours()
    )


# ── Construction des données ──────────────────────────────────────────────────

def _build_conseil_data(class_group, semester):
    """
    Retourne (ues_headers, rows).
    Priorité : données des Bulletins sauvegardés (BulletinUEResult / BulletinECResult).
    Fallback : compute_bulletin_data() en live si aucun bulletin généré.
    """
    from .models import Bulletin, BulletinUEResult, BulletinECResult
    from .services import compute_bulletin_data, _mention
    from academic_core.apps.students.models import Enrollment

    enrollments = Enrollment.objects.filter(
        class_group=class_group,
        status=Enrollment.STATUS_VALIDATED,
    ).select_related('student__user').order_by(
        'student__user__last_name', 'student__user__first_name'
    )
    students = [e.student for e in enrollments]

    # ── Structure UE/EC depuis la maquette ─────────────────────────────────────
    program = class_group.program
    ues = UniteEnseignement.objects.filter(
        program=program, semester=semester
    ).prefetch_related('subjects').order_by('order', 'code')

    ues_headers = []
    for ue in ues:
        subjects = sorted(ue.subjects.all(), key=lambda s: natural_sort_key(s.code))
        if subjects:
            ues_headers.append({'ue': ue, 'subjects': subjects})

    # ── Bulletins sauvegardés ─────────────────────────────────────────────────
    bulletins = Bulletin.objects.filter(
        student__in=students, semester=semester
    ).prefetch_related(
        'ue_results__ue',
        'ue_results__ec_results__subject',
    )
    bulletin_map = {b.student_id: b for b in bulletins}

    rows = []
    for enr in enrollments:
        student = enr.student
        bul = bulletin_map.get(student.pk)

        if bul:
            # ── Source : bulletin sauvegardé ──────────────────────────────────
            ue_res_map = {ur.ue_id: ur for ur in bul.ue_results.all()}
            ue_results = []
            for ue_hdr in ues_headers:
                ur = ue_res_map.get(ue_hdr['ue'].pk)
                if ur:
                    ec_res_map = {ecr.subject_id: ecr for ecr in ur.ec_results.all()}
                    ec_results = []
                    for subj in ue_hdr['subjects']:
                        ecr = ec_res_map.get(subj.pk)
                        if ecr:
                            ec_results.append({
                                'subject':        subj,
                                'final_average':  ecr.final_average,
                                'cc_average':     ecr.cc_average,
                                'd1_score':       None,
                                'd2_score':       None,
                                'exam_score':     ecr.exam_score,
                                'rattrapage_score': ecr.rattrapage_score,
                                'appreciation':   ecr.appreciation,
                                'is_validated':   ecr.is_validated,
                            })
                        else:
                            ec_results.append(None)
                    ue_results.append({
                        'ue':               ue_hdr['ue'],
                        'average':          ur.average,
                        'is_validated':     ur.is_validated,
                        'credits_obtained': ur.credits_obtained,
                        'credits_possible': sum(s.credits or 0 for s in ue_hdr['subjects']),
                        'ec_results':       ec_results,
                    })
                else:
                    ue_results.append({
                        'ue':               ue_hdr['ue'],
                        'average':          None,
                        'is_validated':     False,
                        'credits_obtained': 0,
                        'credits_possible': sum(s.credits or 0 for s in ue_hdr['subjects']),
                        'ec_results':       [None] * len(ue_hdr['subjects']),
                    })

            moy_v = float(bul.semester_average) if bul.semester_average is not None else None
            total_obt = bul.total_credits_obtained
            total_pos = bul.total_credits_possible
            mention   = bul.mention

        else:
            # ── Fallback : calcul live ────────────────────────────────────────
            data = compute_bulletin_data(student, semester, class_group)
            if data is None:
                continue
            ue_map = {ud['ue'].pk: ud for ud in data.get('ue_list', [])}
            ue_results = []
            for ue_hdr in ues_headers:
                ud = ue_map.get(ue_hdr['ue'].pk)
                if ud:
                    ec_map = {ec['subject'].pk: ec for ec in ud['ec_list']}
                    ue_results.append({
                        'ue':               ue_hdr['ue'],
                        'average':          ud['average'],
                        'is_validated':     ud['is_validated'],
                        'credits_obtained': ud.get('credits_obtained', 0),
                        'credits_possible': sum(s.credits or 0 for s in ue_hdr['subjects']),
                        'ec_results':       [ec_map.get(s.pk) for s in ue_hdr['subjects']],
                    })
                else:
                    ue_results.append({
                        'ue':               ue_hdr['ue'],
                        'average':          None,
                        'is_validated':     False,
                        'credits_obtained': 0,
                        'credits_possible': sum(s.credits or 0 for s in ue_hdr['subjects']),
                        'ec_results':       [None] * len(ue_hdr['subjects']),
                    })
            moy_s   = data.get('semester_average')
            moy_v   = float(moy_s) if moy_s is not None else None
            total_obt = data.get('total_credits_obtained', 0)
            total_pos = data.get('total_credits_possible', 0)
            mention   = data.get('mention', '')

        all_ues_validated = ue_results and all(ur['is_validated'] for ur in ue_results)
        decision = 'Validé' if all_ues_validated else 'Rattrapage'
        rows.append({
            'student':                student,
            'matricule':              student.matricule,
            'full_name':              student.full_name,
            'semester_average':       moy_v,
            'total_credits_obtained': total_obt,
            'total_credits_possible': total_pos,
            'mention':                mention,
            'decision':               decision,
            'ue_results':             ue_results,
        })

    # Rang + tri final par mérite
    ranked = sorted([r for r in rows if r['semester_average'] is not None],
                    key=lambda r: r['semester_average'], reverse=True)
    for i, r in enumerate(ranked, 1):
        r['rank'] = i
    for r in rows:
        r.setdefault('rank', 9999)

    rows = sorted(rows, key=lambda r: r['rank'])

    return ues_headers, rows


# ── Vues ──────────────────────────────────────────────────────────────────────

@login_required
def conseil_list(request):
    if not _can_access(request.user):
        messages.error(request, "Accès non autorisé.")
        return redirect('dashboard:index')

    from academic_core.apps.academic_structure.models import AcademicYear

    dept    = getattr(request, 'active_department', None)
    faculty = getattr(request, 'active_faculty', None)
    fac_for_years = faculty or (dept.faculty if dept else None)

    classes_qs = Class.objects.select_related('program__department', 'level').order_by('name')
    if dept:
        classes_qs = classes_qs.filter(program__department=dept)
    elif faculty:
        classes_qs = classes_qs.filter(program__department__faculty=faculty)

    # Année académique : par défaut, seules les classes de l'année en cours
    # sont proposées ici — celles d'une année précédente ne doivent pas rester
    # mélangées avec l'année en cours (même principe que Classes/Maquettes).
    academic_years = AcademicYear.objects.all().order_by('-start_date')
    academic_years = academic_years.filter(faculty=fac_for_years) if fac_for_years else academic_years.none()
    year_param = request.GET.get('year')
    if year_param == 'all':
        selected_year = None
    elif year_param:
        selected_year = academic_years.filter(pk=year_param).first()
        classes_qs = classes_qs.filter(academic_year=selected_year) if selected_year else classes_qs.none()
    else:
        selected_year = academic_years.filter(is_current=True).first()
        if selected_year:
            classes_qs = classes_qs.filter(academic_year=selected_year)

    semesters_qs = Semester.objects.select_related('academic_year').order_by(
        '-academic_year__start_date', 'number'
    )
    if faculty:
        semesters_qs = semesters_qs.filter(academic_year__faculty=faculty)
    if selected_year:
        semesters_qs = semesters_qs.filter(academic_year=selected_year)

    # Si paramètres GET → afficher le tableau directement
    class_id    = request.GET.get('class_id')
    semester_id = request.GET.get('semester_id')

    if class_id and semester_id:
        return redirect('grades:conseil_view',
                        class_id=int(class_id), semester_id=int(semester_id))

    return render(request, 'grades/conseil_list.html', {
        'academic_years': academic_years,
        'selected_year_param': year_param or '',
        'classes':   classes_qs,
        'semesters': semesters_qs,
    })


def _build_annual_recap(class_group, semester_s2, rows_s2):
    """
    Construit le récap annuel quand on est au S2.
    Cherche le S1 de la même année académique et combine les crédits.
    Décision :
      - 60/60 crédits  → Année Validée
      - 42–59 crédits  → Passage Conditionnel
      - < 42  crédits  → Autorisé à Redoubler
    """
    try:
        semester_s1 = Semester.objects.get(
            number=1,
            academic_year=semester_s2.academic_year,
        )
    except Semester.DoesNotExist:
        return None, None

    _, rows_s1 = _build_conseil_data(class_group, semester_s1)
    s1_map = {r['student'].pk: r for r in rows_s1}

    annual_rows = []
    for row_s2 in rows_s2:
        sid = row_s2['student'].pk
        row_s1 = s1_map.get(sid)

        cred_s1_obt = row_s1['total_credits_obtained'] if row_s1 else 0
        cred_s1_pos = row_s1['total_credits_possible'] if row_s1 else 0
        cred_s2_obt = row_s2['total_credits_obtained']
        cred_s2_pos = row_s2['total_credits_possible']

        total_obt = cred_s1_obt + cred_s2_obt
        total_pos = cred_s1_pos + cred_s2_pos   # normalement 60

        moy_s1 = row_s1['semester_average'] if row_s1 else None
        moy_s2 = row_s2['semester_average']

        # Décision annuelle
        if total_obt >= total_pos:
            annual_decision = 'Année Validée'
        elif total_obt >= 42:
            annual_decision = 'Passage Conditionnel'
        else:
            annual_decision = 'Autorisé à Redoubler'

        annual_rows.append({
            'student':     row_s2['student'],
            'matricule':   row_s2['matricule'],
            'full_name':   row_s2['full_name'],
            'moy_s1':      moy_s1,
            'moy_s2':      moy_s2,
            'dec_s1':      row_s1['decision'] if row_s1 else '—',
            'dec_s2':      row_s2['decision'],
            'cred_s1':     f"{cred_s1_obt}/{cred_s1_pos}",
            'cred_s2':     f"{cred_s2_obt}/{cred_s2_pos}",
            'total_obt':   total_obt,
            'total_pos':   total_pos,
            'annual_decision': annual_decision,
        })

    # Rang annuel (sur moy des deux semestres)
    def avg_annuelle(r):
        avgs = [v for v in [r['moy_s1'], r['moy_s2']] if v is not None]
        return sum(avgs) / len(avgs) if avgs else 0

    ranked = sorted(annual_rows, key=avg_annuelle, reverse=True)
    for i, r in enumerate(ranked, 1):
        r['rank'] = i

    return semester_s1, annual_rows


@login_required
def conseil_view(request, class_id, semester_id):
    if not _can_access(request.user):
        messages.error(request, "Accès non autorisé.")
        return redirect('dashboard:index')

    class_group = get_object_or_404(Class, pk=class_id)
    semester    = get_object_or_404(Semester, pk=semester_id)
    ues_headers, rows = _build_conseil_data(class_group, semester)

    # Statistiques semestrielles (rows déjà triés par rang dans _build_conseil_data)
    total       = len(rows)
    valides     = [r for r in rows if r['decision'] == 'Validé']
    rattrapages = [r for r in rows if r['decision'] == 'Rattrapage']
    nb_valides  = len(valides)
    nb_ratt     = len(rattrapages)
    # Rang local dans chaque sous-liste
    for i, r in enumerate(valides, 1):
        r['local_rank'] = i
    for i, r in enumerate(rattrapages, 1):
        r['local_rank'] = i

    # Récap annuel (uniquement pour le semestre 2)
    annual_rows = None
    semester_s1 = None
    if semester.number == 2:
        semester_s1, annual_rows = _build_annual_recap(class_group, semester, rows)

    # Statistiques par UE
    ue_stats = []
    for i, ue_hdr in enumerate(ues_headers):
        nb_val = 0
        nb_rat = 0
        for r in rows:
            if i < len(r['ue_results']):
                ur = r['ue_results'][i]
                if ur and ur.get('average') is not None:
                    if ur['is_validated']:
                        nb_val += 1
                    else:
                        nb_rat += 1
        credits_possible = sum(s.credits or 0 for s in ue_hdr['subjects'])
        ue_stats.append({
            'ue':            ue_hdr['ue'],
            'nb_valide':     nb_val,
            'nb_rattrapage': nb_rat,
            'pct_valide':    round(nb_val / total * 100, 1) if total else 0,
            'credits':       credits_possible,
        })

    return render(request, 'grades/conseil_view.html', {
        'class_group':   class_group,
        'semester':      semester,
        'ues_headers':   ues_headers,
        'rows':          rows,
        'class_id':      class_id,
        'semester_id':   semester_id,
        'nb_total':      total,
        'nb_valides':    nb_valides,
        'nb_rattrapages': nb_ratt,
        'taux_validation': round(nb_valides / total * 100, 1) if total else 0,
        'taux_rattrapage': round(nb_ratt   / total * 100, 1) if total else 0,
        'valides':       valides,
        'rattrapages':   rattrapages,
        # Annuel
        'annual_rows':   annual_rows,
        'semester_s1':   semester_s1,
        # Par UE
        'ue_stats':      ue_stats,
    })


# ── Export Excel ──────────────────────────────────────────────────────────────

@login_required
def conseil_export_excel(request, class_id, semester_id):
    if not _can_access(request.user):
        return redirect('dashboard:index')

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, "openpyxl non installé.")
        return redirect('grades:conseil_view', class_id=class_id, semester_id=semester_id)

    class_group = get_object_or_404(Class, pk=class_id)
    semester    = get_object_or_404(Semester, pk=semester_id)
    ues_headers, rows = _build_conseil_data(class_group, semester)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tableau de conseil"

    NAVY  = "1E3A5F"
    BLUE  = "1E40AF"
    BLUE2 = "3B82F6"
    GREEN_BG = "DCFCE7"
    RED_BG   = "FEE2E2"
    BLUE_BG  = "DBEAFE"
    GRAY     = "F8FAFC"
    WHITE    = "FFFFFF"
    GREEN    = "166534"
    RED      = "DC2626"

    thin = Side(style='thin', color='CBD5E0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _s(cell, text='', bold=False, bg=None, fg='000000', center=True, size=9, wrap=False):
        cell.value = text
        cell.font = Font(bold=bold, color=fg, size=size)
        cell.alignment = Alignment(
            horizontal='center' if center else 'left',
            vertical='center', wrap_text=wrap
        )
        if bg:
            cell.fill = PatternFill('solid', fgColor=bg)
        cell.border = border

    # ── Titre ─────────────────────────────────────────────────────────────────
    n_cols = 3 + sum(len(h['subjects']) + 1 for h in ues_headers) + 5
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    _s(ws.cell(1, 1),
       f"TABLEAU DE CONSEIL — {class_group.name} — {semester.label} — {semester.academic_year}",
       bold=True, bg=NAVY, fg=WHITE, size=12)
    ws.row_dimensions[1].height = 28

    # ── Ligne 2 : UE headers ──────────────────────────────────────────────────
    ws.row_dimensions[2].height = 38
    ws.row_dimensions[3].height = 42

    # Colonnes fixes (fusionnées lignes 2-3)
    for c, lbl in [(1, '#'), (2, 'Matricule'), (3, 'Nom complet')]:
        ws.merge_cells(start_row=2, start_column=c, end_row=3, end_column=c)
        _s(ws.cell(2, c), lbl, bold=True, bg=NAVY, fg=WHITE, wrap=True)

    col = 4
    ue_spans = []  # (ue_hdr, col_start, col_end_incl_moy)
    for ue_hdr in ues_headers:
        n = len(ue_hdr['subjects'])
        col_end = col + n  # col + n-1 pour les EC, +1 pour Moy UE → col+n
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col_end)
        _s(ws.cell(2, col),
           f"{ue_hdr['ue'].code} — {ue_hdr['ue'].title}",
           bold=True, bg=BLUE, fg=WHITE, wrap=True)
        ue_spans.append((ue_hdr, col, col_end))
        col = col_end + 1

    col_moy_sem  = col;  col += 1
    col_credits  = col;  col += 1
    col_mention  = col;  col += 1
    col_decision = col;  col += 1
    col_rang     = col

    for c, lbl in [
        (col_moy_sem, 'Moy\nSem.'), (col_credits, 'Crédits'),
        (col_mention, 'Mention'), (col_decision, 'Décision'), (col_rang, 'Rang'),
    ]:
        ws.merge_cells(start_row=2, start_column=c, end_row=3, end_column=c)
        _s(ws.cell(2, c), lbl, bold=True, bg=NAVY, fg=WHITE, wrap=True)

    # ── Ligne 3 : EC headers ──────────────────────────────────────────────────
    for ue_hdr, col_s, col_e in ue_spans:
        c = col_s
        for subj in ue_hdr['subjects']:
            _s(ws.cell(3, c),
               f"{subj.code}\n{subj.title[:20]}\ncoef {subj.coefficient}",
               bold=True, bg=BLUE2, fg=WHITE, wrap=True)
            c += 1
        _s(ws.cell(3, c), 'Moy UE', bold=True, bg=BLUE, fg=WHITE, wrap=True)

    # ── Données ───────────────────────────────────────────────────────────────
    for idx, row in enumerate(rows, 1):
        r = 3 + idx
        ws.row_dimensions[r].height = 16
        row_bg = GRAY if idx % 2 == 0 else WHITE

        _s(ws.cell(r, 1), idx, center=True, bg=row_bg)
        _s(ws.cell(r, 2), row['matricule'] or '', center=True, bg=row_bg)
        _s(ws.cell(r, 3), row['full_name'], center=False, bg=row_bg)

        for ue_hdr, col_s, col_e in ue_spans:
            ue_res = next((u for u in row['ue_results'] if u['ue'].pk == ue_hdr['ue'].pk), None)
            c = col_s
            if ue_res:
                for ec in ue_res['ec_results']:
                    if ec and ec['final_average'] is not None:
                        v = float(ec['final_average'])
                        _s(ws.cell(r, c), round(v, 2), center=True,
                           fg=GREEN if v >= 10 else RED, bg=row_bg)
                    else:
                        _s(ws.cell(r, c), '—', center=True, bg=row_bg, fg='94A3B8')
                    c += 1
                moy_u = ue_res['average']
                vu = float(moy_u) if moy_u is not None else None
                _s(ws.cell(r, c), round(vu, 2) if vu is not None else '—',
                   bold=True, center=True,
                   fg=GREEN if (vu and vu >= 10) else RED,
                   bg=BLUE_BG if (vu and vu >= 10) else RED_BG)
            else:
                for i in range(col_e - col_s + 1):
                    _s(ws.cell(r, col_s + i), '—', center=True, bg=row_bg, fg='94A3B8')

        moy_v = row['semester_average']
        dec   = row['decision']
        _s(ws.cell(r, col_moy_sem),
           round(moy_v, 2) if moy_v else '—', bold=True, center=True,
           fg=GREEN if (moy_v and moy_v >= 10) else RED,
           bg=GREEN_BG if (moy_v and moy_v >= 10) else RED_BG)
        _s(ws.cell(r, col_credits),
           f"{row['total_credits_obtained']}/{row['total_credits_possible']}",
           center=True, bg=row_bg)
        _s(ws.cell(r, col_mention), row['mention'] or '', center=True, bg=row_bg)
        _s(ws.cell(r, col_decision), dec, bold=True, center=True,
           fg=GREEN if dec == 'Validé' else RED,
           bg=GREEN_BG if dec == 'Validé' else RED_BG)
        _s(ws.cell(r, col_rang), '—' if row['rank'] == 9999 else str(row['rank']), bold=True, center=True, bg=row_bg)

    # Largeurs
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 13
    ws.column_dimensions['C'].width = 25
    for i in range(4, col_rang + 1):
        ws.column_dimensions[get_column_letter(i)].width = 10
    ws.column_dimensions[get_column_letter(col_moy_sem)].width = 8
    ws.column_dimensions[get_column_letter(col_mention)].width = 14
    ws.column_dimensions[get_column_letter(col_decision)].width = 12
    ws.freeze_panes = 'D4'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"conseil_{class_group.name}_{semester.label}.xlsx".replace(' ', '_')
    resp = HttpResponse(buf.getvalue(),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{fname}"'
    return resp


# ── Export PDF ────────────────────────────────────────────────────────────────

@login_required
def conseil_export_pdf(request, class_id, semester_id):
    if not _can_access(request.user):
        return redirect('dashboard:index')

    try:
        from reportlab.lib.pagesizes import A3, landscape
        from reportlab.lib import colors
        from reportlab.lib.units import cm, mm
        from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                        Paragraph, Spacer)
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from academic_core.pdf_utils import watermark_canvas, logo_image, get_institut_config_for_request
    except ImportError:
        messages.error(request, "reportlab non installé.")
        return redirect('grades:conseil_view', class_id=class_id, semester_id=semester_id)

    class_group = get_object_or_404(Class, pk=class_id)
    semester    = get_object_or_404(Semester, pk=semester_id)
    ues_headers, rows = _build_conseil_data(class_group, semester)

    MARGIN = 0.8 * cm
    PAGE_W, PAGE_H = landscape(A3)          # 420mm × 297mm
    USABLE_W = PAGE_W - 2 * MARGIN          # ≈ 403mm

    # Largeurs fixes
    W_NUM   = 0.65 * cm
    W_MAT   = 2.3  * cm
    W_NOM   = 4.8  * cm
    W_FIXED = W_NUM + W_MAT + W_NOM

    # Colonnes synthèse (droite)
    W_MOY_S = 1.3 * cm
    W_CRED  = 1.3 * cm
    W_MENT  = 2.2 * cm
    W_DEC   = 1.6 * cm
    W_RANG  = 0.9 * cm
    W_SYNTH = W_MOY_S + W_CRED + W_MENT + W_DEC + W_RANG

    # Colonnes EC dynamiques — répartir le reste
    n_ec_cols = sum(len(h['subjects']) + 1 for h in ues_headers)  # ECs + Moy UE par UE
    W_EC = max(1.1 * cm, (USABLE_W - W_FIXED - W_SYNTH) / n_ec_cols) if n_ec_cols else 1.5 * cm

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A3),
                             leftMargin=MARGIN, rightMargin=MARGIN,
                             topMargin=1.2*cm, bottomMargin=0.8*cm)

    navy_c  = colors.HexColor('#1E3A5F')
    blue_c  = colors.HexColor('#1E40AF')
    blue2_c = colors.HexColor('#3B82F6')
    green_c = colors.HexColor('#166534')
    red_c   = colors.HexColor('#DC2626')
    oran_c  = colors.HexColor('#C2410C')
    gray_c  = colors.HexColor('#F8FAFC')
    gbg_c   = colors.HexColor('#DCFCE7')
    rbg_c   = colors.HexColor('#FFF7ED')

    FS = 6.5   # font size EC data
    FH = 6.5   # font size header

    def P(text, bold=False, size=FS, color=colors.black, align=TA_CENTER):
        st = ParagraphStyle('_', fontSize=size, textColor=color,
                             fontName='Helvetica-Bold' if bold else 'Helvetica',
                             alignment=align, leading=size + 1.2, wordWrap='CJK')
        return Paragraph(str(text) if text is not None else '—', st)

    # ── Construction des lignes d'en-tête ─────────────────────────────────────
    h1 = [P('#', True, FH, colors.white),
          P('Matricule', True, FH, colors.white),
          P('Nom complet', True, FH, colors.white, TA_LEFT)]
    h2 = ['', '', '']
    col_w = [W_NUM, W_MAT, W_NOM]
    spans = [('SPAN',(0,0),(0,1)), ('SPAN',(1,0),(1,1)), ('SPAN',(2,0),(2,1))]
    ci = 3

    for ue_hdr in ues_headers:
        n = len(ue_hdr['subjects'])
        label = f"{ue_hdr['ue'].code}\n{ue_hdr['ue'].title[:30]}"
        h1.append(P(label, True, FH, colors.white))
        h1 += [''] * n
        h2 += [P(f"{s.code}\nc{s.coefficient}", True, FH - 0.5, colors.white)
               for s in ue_hdr['subjects']]
        h2.append(P("Moy\nUE", True, FH, colors.white))
        col_w += [W_EC] * n + [W_EC * 1.1]
        spans.append(('SPAN', (ci, 0), (ci + n, 0)))
        ci += n + 1

    for lbl in ['Moy\nSem.', 'Crédits', 'Mention', 'Décision', 'Rang']:
        h1.append(P(lbl, True, FH, colors.white))
        h2.append('')
        spans.append(('SPAN', (ci, 0), (ci, 1)))
        ci += 1
    col_w += [W_MOY_S, W_CRED, W_MENT, W_DEC, W_RANG]

    tdata = [h1, h2]

    # ── Lignes données ─────────────────────────────────────────────────────────
    for idx, row in enumerate(rows, 1):
        bg_row = colors.white if idx % 2 == 1 else gray_c
        cells = [
            P(str(idx), size=FS),
            P(row['matricule'] or '', size=FS),
            P(row['full_name'], size=FS, align=TA_LEFT),
        ]
        for ue_hdr in ues_headers:
            ur = next((u for u in row['ue_results'] if u['ue'].pk == ue_hdr['ue'].pk), None)
            if ur:
                for ec in ur['ec_results']:
                    if ec and ec['final_average'] is not None:
                        v = float(ec['final_average'])
                        cells.append(P(f"{v:.2f}", size=FS, color=green_c if v >= 10 else red_c))
                    else:
                        cells.append(P('—', size=FS, color=colors.grey))
                vu = float(ur['average']) if ur['average'] is not None else None
                cells.append(P(f"{vu:.2f}" if vu else '—', True, FS,
                               green_c if (vu and vu >= 10) else red_c))
            else:
                cells += [P('—', size=FS, color=colors.grey)] * (len(ue_hdr['subjects']) + 1)

        mv  = row['semester_average']
        dec = row['decision']
        ok  = dec == 'Validé'
        cells += [
            P(f"{mv:.2f}" if mv else '—', True, FS + 0.5, green_c if (mv and mv >= 10) else red_c),
            P(f"{row['total_credits_obtained']}/{row['total_credits_possible']}", size=FS),
            P(row['mention'] or '—', size=FS),
            P(dec, True, FS, green_c if ok else oran_c),
            P('—' if row['rank'] == 9999 else str(row['rank']), True, FS),
        ]
        tdata.append(cells)

    # ── Style ─────────────────────────────────────────────────────────────────
    ts = TableStyle([
        ('BACKGROUND', (0,0), (-1,0), navy_c),
        ('BACKGROUND', (0,1), (-1,1), blue2_c),
        ('GRID',       (0,0), (-1,-1), 0.25, colors.HexColor('#CBD5E0')),
        ('VALIGN',     (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2),
        ('BOTTOMPADDING',(0,0),(-1,-1), 2),
        ('LEFTPADDING', (0,0), (-1,-1), 2),
        ('RIGHTPADDING',(0,0), (-1,-1), 2),
        ('ROWBACKGROUNDS', (0,2), (-1,-1), [colors.white, gray_c]),
    ] + spans)

    # Couleurs moy UE et colonnes synthèse par ligne
    for i, row in enumerate(rows, 2):
        col_cur = 3
        for ue_hdr in ues_headers:
            n = len(ue_hdr['subjects'])
            col_moy_ue = col_cur + n
            ur = next((u for u in row['ue_results'] if u['ue'].pk == ue_hdr['ue'].pk), None)
            if ur and ur['is_validated']:
                ts.add('BACKGROUND', (col_moy_ue, i), (col_moy_ue, i), colors.HexColor('#DBEAFE'))
            elif ur and ur['average'] is not None:
                ts.add('BACKGROUND', (col_moy_ue, i), (col_moy_ue, i), colors.HexColor('#FEE2E2'))
            col_cur += n + 1
        # Décision
        dec = row['decision']
        ok  = dec == 'Validé'
        moy_col = ci - 5
        dec_col = ci - 2
        ts.add('BACKGROUND', (moy_col, i), (moy_col, i), gbg_c if (row['semester_average'] and row['semester_average'] >= 10) else rbg_c)
        ts.add('BACKGROUND', (dec_col, i), (dec_col, i), gbg_c if ok else rbg_c)

    t = Table(tdata, colWidths=col_w, repeatRows=2)
    t.setStyle(ts)

    title_st = ParagraphStyle('T', fontSize=12, fontName='Helvetica-Bold',
                               textColor=navy_c, alignment=TA_CENTER, spaceAfter=3)
    sub_st   = ParagraphStyle('S', fontSize=8, textColor=colors.HexColor('#64748B'),
                               alignment=TA_CENTER, spaceAfter=8)

    config = get_institut_config_for_request(request)
    _logo = logo_image(width=1.6*cm, height=1.6*cm, config=config)
    inst_name = (config.nom if config else '') or "Institut Supérieur d'Informatique - ISI"
    logo_row = Table([[
        _logo or '',
        Paragraph(f'<b>{inst_name}</b>', ParagraphStyle(
            'InstName', fontSize=10, fontName='Helvetica-Bold',
            textColor=navy_c, alignment=TA_CENTER,
        )),
    ]], colWidths=[2.2*cm, None])
    logo_row.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))

    from reportlab.platypus import HRFlowable
    story = [
        logo_row,
        HRFlowable(width='100%', thickness=1, color=navy_c, spaceAfter=4),
        Paragraph("TABLEAU DE CONSEIL", title_st),
        Paragraph(f"{class_group.name}  ·  {semester.label}  ·  {semester.academic_year}", sub_st),
        t,
    ]
    doc.build(story, onFirstPage=watermark_canvas, onLaterPages=watermark_canvas)
    buf.seek(0)
    fname = f"conseil_{class_group.name}_{semester.label}.pdf".replace(' ', '_')
    resp = HttpResponse(buf.getvalue(), content_type='application/pdf')
    resp['Content-Disposition'] = f'attachment; filename="{fname}"'
    return resp


# ── Export Word ───────────────────────────────────────────────────────────────

@login_required
def conseil_export_word(request, class_id, semester_id):
    if not _can_access(request.user):
        return redirect('dashboard:index')

    try:
        import os
        from docx import Document
        from docx.shared import Pt, RGBColor, Cm, Twips
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_ALIGN_VERTICAL
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        from docx.enum.section import WD_ORIENT
        from academic_core.pdf_utils import add_docx_watermark, get_logo_path, get_institut_config_for_request
    except ImportError:
        messages.error(request, "python-docx non installé.")
        return redirect('grades:conseil_view', class_id=class_id, semester_id=semester_id)

    class_group = get_object_or_404(Class, pk=class_id)
    semester    = get_object_or_404(Semester, pk=semester_id)
    ues_headers, rows = _build_conseil_data(class_group, semester)

    # ── Helpers ───────────────────────────────────────────────────────────────
    def rgb(h):
        h = h.lstrip('#')
        return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

    def set_cell_shd(cell, hex_fill):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        # retire ancien shd
        for old in tcPr.findall(qn('w:shd')):
            tcPr.remove(old)
        s = OxmlElement('w:shd')
        s.set(qn('w:val'),   'clear')
        s.set(qn('w:color'), 'auto')
        s.set(qn('w:fill'),  hex_fill.lstrip('#'))
        tcPr.append(s)

    def set_col_w(cell, width_cm):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        for old in tcPr.findall(qn('w:tcW')):
            tcPr.remove(old)
        tcW = OxmlElement('w:tcW')
        tcW.set(qn('w:w'),    str(int(width_cm * 567)))  # 1cm ≈ 567 twips
        tcW.set(qn('w:type'), 'dxa')
        tcPr.append(tcW)

    def wc(cell, text, bold=False, size=7, color='000000', align='center', fill=None):
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == 'center' else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(str(text) if text is not None else '—')
        run.bold      = bold
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
        if fill:
            set_cell_shd(cell, fill)

    # ── Page A3 paysage ───────────────────────────────────────────────────────
    doc = Document()
    sec = doc.sections[0]
    # A3 = 297mm × 420mm → paysage = 420mm × 297mm
    sec.orientation  = WD_ORIENT.LANDSCAPE
    sec.page_width   = Cm(42)
    sec.page_height  = Cm(29.7)
    sec.left_margin  = sec.right_margin  = Cm(0.7)
    sec.top_margin   = sec.bottom_margin = Cm(1.0)

    # Largeur utile ≈ 42 - 1.4 = 40.6 cm
    USABLE = 40.6
    W_NUM, W_MAT, W_NOM = 0.6, 2.1, 4.5
    W_MOY_S, W_CRED, W_MENT, W_DEC, W_RANG = 1.2, 1.3, 2.0, 1.8, 0.8
    n_ec_cols = sum(len(h['subjects']) + 1 for h in ues_headers)
    W_EC_RAW  = (USABLE - W_NUM - W_MAT - W_NOM - W_MOY_S - W_CRED - W_MENT - W_DEC - W_RANG) / max(n_ec_cols, 1)
    W_EC      = max(1.05, W_EC_RAW)
    W_MOY_UE  = W_EC * 1.1

    # ── En-tête : logo institut + nom ─────────────────────────────────────────
    config = get_institut_config_for_request(request)
    inst_nom = (getattr(config, 'nom', None)) or "Institut Supérieur d'Informatique — ISI"
    logo_path = get_logo_path(config)
    hdr = doc.add_table(rows=1, cols=2)
    hdr.autofit = False
    hdr.columns[0].width = Cm(2.2)
    cell_logo = hdr.cell(0, 0)
    cell_logo.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if logo_path and os.path.exists(logo_path):
        p_logo.add_run().add_picture(logo_path, height=Cm(1.4))
    cell_nom = hdr.cell(0, 1)
    cell_nom.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p_nom = cell_nom.paragraphs[0]
    run_nom = p_nom.add_run(inst_nom)
    run_nom.bold = True
    run_nom.font.size = Pt(11)
    run_nom.font.color.rgb = rgb('1E3A5F')
    doc.add_paragraph()

    # ── Titre ────────────────────────────────────────────────────────────────
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tp.add_run(f"TABLEAU DE CONSEIL  —  {class_group.name}  —  {semester.label}  —  {semester.academic_year}")
    run.bold = True; run.font.size = Pt(11); run.font.color.rgb = rgb('1E3A5F')

    # ── Calcul nb colonnes ────────────────────────────────────────────────────
    n_ec_total = sum(len(h['subjects']) for h in ues_headers)
    n_cols = 3 + n_ec_total + len(ues_headers) + 5

    table = doc.add_table(rows=2 + len(rows), cols=n_cols)
    table.style = 'Table Grid'

    # ── Ligne 0 : UE headers ─────────────────────────────────────────────────
    r0 = table.rows[0]
    for c, lbl, ww in [(0,'#',W_NUM),(1,'Matricule',W_MAT),(2,'Nom complet',W_NOM)]:
        wc(r0.cells[c], lbl, True, 7, 'FFFFFF', fill='1E3A5F')
        set_col_w(r0.cells[c], ww)

    ci = 3
    for ue_hdr in ues_headers:
        n = len(ue_hdr['subjects'])
        merged = r0.cells[ci]
        for x in range(1, n + 1):
            merged.merge(r0.cells[ci + x])
        wc(merged, f"{ue_hdr['ue'].code}  –  {ue_hdr['ue'].title}", True, 7, 'FFFFFF', fill='1E40AF')
        ci += n + 1

    for lbl, ww in [('Moy\nSem.',W_MOY_S),('Crédits',W_CRED),('Mention',W_MENT),('Décision',W_DEC),('Rang',W_RANG)]:
        wc(r0.cells[ci], lbl, True, 7, 'FFFFFF', fill='1E3A5F')
        set_col_w(r0.cells[ci], ww)
        ci += 1

    # ── Ligne 1 : EC headers ──────────────────────────────────────────────────
    r1 = table.rows[1]
    for c, lbl, ww in [(0,'#',W_NUM),(1,'Matricule',W_MAT),(2,'Nom complet',W_NOM)]:
        wc(r1.cells[c], lbl, True, 6, 'FFFFFF', fill='1E3A5F')
        set_col_w(r1.cells[c], ww)
    ci = 3
    for ue_hdr in ues_headers:
        for subj in ue_hdr['subjects']:
            wc(r1.cells[ci], f"{subj.code}\nc{subj.coefficient}", True, 6, 'FFFFFF', fill='3B82F6')
            set_col_w(r1.cells[ci], W_EC)
            ci += 1
        wc(r1.cells[ci], 'Moy\nUE', True, 6, 'FFFFFF', fill='1D4ED8')
        set_col_w(r1.cells[ci], W_MOY_UE)
        ci += 1

    # ── Lignes données ────────────────────────────────────────────────────────
    for idx, row in enumerate(rows):
        tr  = table.rows[2 + idx]
        bg  = 'F8FAFC' if idx % 2 == 0 else 'FFFFFF'
        wc(tr.cells[0], str(idx+1), size=7, fill=bg)
        wc(tr.cells[1], row['matricule'] or '', size=7, fill=bg)
        wc(tr.cells[2], row['full_name'], size=7, align='left', fill=bg)

        ci = 3
        for ue_hdr in ues_headers:
            ur = next((u for u in row['ue_results'] if u['ue'].pk == ue_hdr['ue'].pk), None)
            if ur:
                for ec in ur['ec_results']:
                    if ec and ec['final_average'] is not None:
                        v   = float(ec['final_average'])
                        clr = '166534' if v >= 10 else 'DC2626'
                        wc(tr.cells[ci], f"{v:.2f}", size=7, color=clr, fill=bg)
                    else:
                        wc(tr.cells[ci], '—', size=7, color='94A3B8', fill=bg)
                    ci += 1
                vu = float(ur['average']) if ur['average'] else None
                vc = '166534' if (vu and vu >= 10) else 'DC2626'
                vb = 'DBEAFE' if (vu and vu >= 10) else 'FEE2E2'
                wc(tr.cells[ci], f"{vu:.2f}" if vu else '—', True, 7, vc, fill=vb)
                ci += 1
            else:
                for _ in range(len(ue_hdr['subjects']) + 1):
                    wc(tr.cells[ci], '—', size=7, color='94A3B8', fill=bg); ci += 1

        mv  = row['semester_average']
        dec = row['decision']
        ok  = dec == 'Validé'
        dc  = '166534' if ok else 'C2410C'
        db  = 'DCFCE7' if ok else 'FFF7ED'
        mb  = 'DCFCE7' if (mv and mv >= 10) else 'FEE2E2'
        mc  = '166534' if (mv and mv >= 10) else 'DC2626'
        for val, color, fill, bld in [
            (f"{mv:.2f}" if mv else '—', mc, mb, True),
            (f"{row['total_credits_obtained']}/{row['total_credits_possible']}", '000000', bg, False),
            (row['mention'] or '—', '1E3A5F', bg, False),
            (dec, dc, db, True),
            ('—' if row['rank'] == 9999 else str(row['rank']), '1E3A5F', bg, True),
        ]:
            wc(tr.cells[ci], val, bld, 7, color, fill=fill); ci += 1

    # ── Hauteur minimale des lignes ───────────────────────────────────────────
    for row_obj in table.rows:
        trPr = row_obj._tr.get_or_add_trPr()
        trH  = OxmlElement('w:trHeight')
        trH.set(qn('w:val'), '280')   # ≈ 0.5 cm
        trH.set(qn('w:hRule'), 'atLeast')
        trPr.append(trH)

    add_docx_watermark(doc)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    fname = f"conseil_{class_group.name}_{semester.label}.docx".replace(' ', '_')
    resp = HttpResponse(buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    resp['Content-Disposition'] = f'attachment; filename="{fname}"'
    return resp
