"""
Générateur de rapports Excel avec openpyxl.
"""
import io
import os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

PRIMARY_HEX = '0D6EFD'
LIGHT_HEX   = 'F8F9FA'
SUCCESS_HEX = '198754'
DANGER_HEX  = 'DC3545'


def _header_style():
    return Font(bold=True, color='FFFFFF', size=10)

def _header_fill():
    return PatternFill('solid', fgColor=PRIMARY_HEX)

def _light_fill():
    return PatternFill('solid', fgColor=LIGHT_HEX)

def _thin_border():
    thin = Side(style='thin', color='DEE2E6')
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def _auto_col_width(ws):
    for col in ws.columns:
        max_w = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_w + 4, 50)


def _add_logo_to_sheet(ws, config=None, cell='A1', height_px=45):
    """Insère le logo de l'institut dans la feuille Excel à la cellule indiquée."""
    from academic_core.pdf_utils import get_logo_path
    logo_path = get_logo_path(config)
    if logo_path and os.path.exists(logo_path):
        try:
            img = XLImage(logo_path)
            ratio = img.width / img.height if img.height else 1
            img.height = height_px
            img.width  = int(height_px * ratio)
            ws.add_image(img, cell)
        except Exception:
            pass


def _inst_header_row(ws, config, row_num, merge_to_col):
    """Insère une ligne d'en-tête avec le nom de l'institut."""
    inst_nom = (config.nom if config else None) or "Institut Supérieur d'Informatique - ISI"
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=merge_to_col)
    cell = ws.cell(row=row_num, column=1, value=inst_nom.upper())
    cell.font      = Font(bold=True, size=11, color=PRIMARY_HEX)
    cell.alignment = Alignment(horizontal='center')


def generate_grades_excel(class_group, semester, config=None):
    from academic_core.apps.grades.models import Grade, SubjectAverage
    from academic_core.apps.students.models import Enrollment

    wb = Workbook()
    ws = wb.active
    ws.title = f"Notes {semester.label}"

    # Logo en A1
    _add_logo_to_sheet(ws, config=config, cell='A1', height_px=40)
    ws.row_dimensions[1].height = 32

    # Nom de l'institut
    _inst_header_row(ws, config, row_num=2, merge_to_col=8)

    # Titre du document
    ws.merge_cells('A3:H3')
    ws['A3'] = f"Relevé de Notes — {class_group.name} — {semester.label}"
    ws['A3'].font      = Font(bold=True, size=13, color=PRIMARY_HEX)
    ws['A3'].alignment = Alignment(horizontal='center')
    ws['A4'] = f"Généré le {date.today().strftime('%d/%m/%Y')}"
    ws['A4'].font = Font(italic=True, size=9, color='6C757D')

    from academic_core.apps.subjects.models import natural_sort_key
    subjects  = sorted(
        semester.subjects.filter(program=class_group.program),
        key=lambda s: natural_sort_key(s.code),
    )
    headers   = ['Matricule', 'Nom & Prénom'] + [s.code for s in subjects] + ['Moyenne', 'Rang']
    row_idx   = 6

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col, value=h)
        cell.font      = _header_style()
        cell.fill      = _header_fill()
        cell.alignment = Alignment(horizontal='center')
        cell.border    = _thin_border()

    enrollments = Enrollment.objects.filter(
        class_group=class_group, academic_year=semester.academic_year,
        status=Enrollment.STATUS_VALIDATED,
    ).select_related('student__user').order_by('student__user__last_name')

    for i, enr in enumerate(enrollments):
        student  = enr.student
        row_idx += 1
        fill     = _light_fill() if i % 2 == 0 else PatternFill('solid', fgColor='FFFFFF')

        ws.cell(row=row_idx, column=1, value=student.matricule).fill = fill
        ws.cell(row=row_idx, column=2, value=student.full_name).fill = fill

        for col, subject in enumerate(subjects, 3):
            sa = SubjectAverage.objects.filter(
                student=student, subject=subject, semester=semester
            ).first()
            val  = float(sa.average) if sa else None
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill   = fill
            cell.number_format = '0.00'
            cell.alignment = Alignment(horizontal='center')
            if val is not None:
                cell.font = Font(
                    color=SUCCESS_HEX if val >= 10 else DANGER_HEX,
                    bold=(val < 10)
                )

        # Semester avg
        from academic_core.apps.grades.models import SemesterAverage
        sem_avg = SemesterAverage.objects.filter(student=student, semester=semester).first()
        avg_col = len(subjects) + 3
        avg_cell = ws.cell(row=row_idx, column=avg_col, value=float(sem_avg.average) if sem_avg else None)
        avg_cell.number_format = '0.00'
        avg_cell.font = Font(bold=True, color=SUCCESS_HEX if (sem_avg and sem_avg.average >= 10) else DANGER_HEX)
        avg_cell.fill = fill

        rank_cell = ws.cell(row=row_idx, column=avg_col + 1, value=sem_avg.rank if sem_avg else None)
        rank_cell.alignment = Alignment(horizontal='center')
        rank_cell.fill = fill

    _auto_col_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_timetable_excel(semester, class_group=None, teacher=None, config=None):
    from academic_core.apps.timetable.models import TimetableEntry

    wb = Workbook()
    ws = wb.active
    ws.title = "Emploi du temps"

    _add_logo_to_sheet(ws, config=config, cell='A1', height_px=40)
    ws.row_dimensions[1].height = 32
    _inst_header_row(ws, config, row_num=2, merge_to_col=9)

    qs = TimetableEntry.objects.filter(semester=semester, is_active=True).select_related(
        'subject', 'teacher__user', 'class_group', 'room'
    ).order_by('day_of_week', 'start_time')
    if class_group:
        qs = qs.filter(class_group=class_group)
    if teacher:
        qs = qs.filter(teacher=teacher)

    headers = ['Jour', 'Début', 'Fin', 'Durée (h)', 'Module (EC)', 'Classe', 'Enseignant', 'Salle', 'Récurrence']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font  = _header_style()
        cell.fill  = _header_fill()
        cell.border = _thin_border()

    for row_i, entry in enumerate(qs, 4):
        fill = _light_fill() if row_i % 2 == 0 else PatternFill('solid', fgColor='FFFFFF')
        data = [
            entry.get_day_of_week_display(),
            str(entry.start_time)[:5],
            str(entry.end_time)[:5],
            round(entry.duration_hours, 2),
            entry.subject.title,
            entry.class_group.name,
            entry.teacher.full_name,
            entry.room.code if entry.room else '—',
            entry.get_recurrence_display(),
        ]
        for col, val in enumerate(data, 1):
            cell = ws.cell(row=row_i, column=col, value=val)
            cell.fill   = fill
            cell.border = _thin_border()

    _auto_col_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_rapport_annuel_excel(detail, config=None):
    """`detail` : dict retourné par grades.annual_report_services.compute_annual_department_report()."""
    department = detail['department']
    academic_year = detail['academic_year']
    s = detail['summary']

    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport Annuel"

    _add_logo_to_sheet(ws, config=config, cell='A1', height_px=40)
    ws.row_dimensions[1].height = 32
    _inst_header_row(ws, config, row_num=2, merge_to_col=5)

    ws.merge_cells('A3:E3')
    ws['A3'] = f"Rapport Annuel de la Direction — {department.name} — {academic_year.label}"
    ws['A3'].font      = Font(bold=True, size=13, color=PRIMARY_HEX)
    ws['A3'].alignment = Alignment(horizontal='center')
    ws['A4'] = f"Généré le {date.today().strftime('%d/%m/%Y')}"
    ws['A4'].font = Font(italic=True, size=9, color='6C757D')

    row_idx = 6

    def _write_bucket_table(title, bucket_dict, start_row):
        r = start_row
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=11, color=PRIMARY_HEX)
        r += 1
        headers = ['Catégorie', 'Effectif', 'Admis', 'Non admis', 'Absent examen']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=r, column=col, value=h)
            cell.font = _header_style()
            cell.fill = _header_fill()
            cell.border = _thin_border()
        r += 1
        for label, row in bucket_dict.items():
            values = [label, row['total'], row['admis'], row['non_admis'], row['non_disponible']]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=r, column=col, value=val)
                cell.border = _thin_border()
            r += 1
        return r + 1

    ws.cell(row=row_idx, column=1, value='Effectif total').font = Font(bold=True)
    ws.cell(row=row_idx, column=2, value=s['total'])
    ws.cell(row=row_idx + 1, column=1, value='Admis').font = Font(bold=True, color=SUCCESS_HEX)
    ws.cell(row=row_idx + 1, column=2, value=s['admis'])
    ws.cell(row=row_idx + 2, column=1, value='Non admis').font = Font(bold=True, color=DANGER_HEX)
    ws.cell(row=row_idx + 2, column=2, value=s['non_admis'])
    ws.cell(row=row_idx + 3, column=1, value='Absent examen').font = Font(bold=True)
    ws.cell(row=row_idx + 3, column=2, value=s['non_disponible'])
    row_idx += 5

    row_idx = _write_bucket_table('Par genre', s['by_gender'], row_idx)
    row_idx = _write_bucket_table('Par nationalité', s['by_nationality'], row_idx)

    ws.cell(row=row_idx, column=1, value='Par classe').font = Font(bold=True, size=11, color=PRIMARY_HEX)
    row_idx += 1
    headers = ['Classe', 'Effectif', 'Admis', 'Non admis', 'Absent examen']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col, value=h)
        cell.font = _header_style()
        cell.fill = _header_fill()
        cell.border = _thin_border()
    row_idx += 1
    for entry in detail['classes']:
        values = [
            entry['class_group'].name, entry['total'], entry['admis'],
            entry['non_admis'], entry['non_disponible'],
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = _thin_border()
        row_idx += 1

    _auto_col_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _class_results_session_groups(session_key):
    """(titre, nb colonnes, rowspan ?, clé du champ dans la session) par
    colonne — colonnes H/F pour les groupes à 2 colonnes (clé -> dict
    {'M':.., 'F':..}), 1 colonne pour les taux/moyenne/effectif (rowspan).
    « Abandons » n'existe que côté Session Normale, « Réclamation » que côté
    Session de Rattrapage (voir class_results_services.py)."""
    common_tail = [
        ('Non Admis', 2, False, 'non_admis'),
        ('Absent Examen', 2, False, 'absent_examen'),
    ]
    if session_key in ('normale', 'normale_reclamation'):
        return (
            [('Classe', 1, True, None), ('Effectif Total', 1, True, 'effectif_total'),
             ('Effectif Composé', 2, False, 'composes'), ('Abandons', 2, False, 'abandons'),
             ('Admis', 2, False, 'admis')]
            + common_tail
            + [('Taux de réussite', 1, True, 'taux_reussite'), ('Taux H', 1, True, 'taux_h'),
               ('Taux F', 1, True, 'taux_f'), ('Meilleure moyenne', 1, True, 'meilleure_moyenne')]
        )
    return (
        [('Classe', 1, True, None), ('Effectif Total', 1, True, 'effectif_total'),
         ('Effectif Composé', 2, False, 'composes'), ('Admis', 2, False, 'admis')]
        + common_tail
        + [('Réclamation', 2, False, 'reclamation'),
           ('Taux de réussite', 1, True, 'taux_reussite'), ('Taux H', 1, True, 'taux_h'),
           ('Taux F', 1, True, 'taux_f'), ('Meilleure moyenne', 1, True, 'meilleure_moyenne')]
    )


def _session_row_values(class_label, session, groups):
    values = [class_label]
    for title, span, is_rowspan, field in groups[1:]:
        if field == 'effectif_total':
            values.append(session['effectif_total'])
        elif field in ('taux_reussite', 'taux_h', 'taux_f'):
            v = session[field]
            values.append(f"{v} %" if v is not None else '—')
        elif field == 'meilleure_moyenne':
            if session['meilleure_moyenne'] is not None:
                values.append(
                    f"{session['meilleure_moyenne']:.2f}/20 — "
                    f"{session['meilleure_moyenne_etudiant'].user.get_full_name()}"
                )
            else:
                values.append('—')
        else:
            values.append(session[field]['M'])
            values.append(session[field]['F'])
    return values


def _write_class_results_sheet(ws, session_key, title_label, rows,
                                config, department, academic_year, semester):
    groups = _class_results_session_groups(session_key)
    total_cols = sum(span for _, span, _, _ in groups)

    _add_logo_to_sheet(ws, config=config, cell='A1', height_px=40)
    ws.row_dimensions[1].height = 32
    _inst_header_row(ws, config, row_num=2, merge_to_col=total_cols)

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=total_cols)
    ws.cell(row=3, column=1, value=(
        f"Résultats par classe — {title_label} — {department.name} — {academic_year.label} — "
        f"{semester.level.name if semester.level else ''} {semester.label}"
    ))
    ws['A3'].font      = Font(bold=True, size=13, color=PRIMARY_HEX)
    ws['A3'].alignment = Alignment(horizontal='center')
    ws['A4'] = f"Généré le {date.today().strftime('%d/%m/%Y')}"
    ws['A4'].font = Font(italic=True, size=9, color='6C757D')

    HEADER_ROW1 = 6
    HEADER_ROW2 = 7

    col = 1
    for title, span, is_rowspan, _field in groups:
        if is_rowspan:
            ws.merge_cells(start_row=HEADER_ROW1, start_column=col, end_row=HEADER_ROW2, end_column=col)
            cell = ws.cell(row=HEADER_ROW1, column=col, value=title)
        else:
            ws.merge_cells(start_row=HEADER_ROW1, start_column=col, end_row=HEADER_ROW1, end_column=col + span - 1)
            cell = ws.cell(row=HEADER_ROW1, column=col, value=title)
            ws.cell(row=HEADER_ROW2, column=col, value='H')
            ws.cell(row=HEADER_ROW2, column=col + 1, value='F')
        cell.font = _header_style()
        cell.fill = _header_fill()
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        col += span

    for c in range(1, total_cols + 1):
        for r in (HEADER_ROW1, HEADER_ROW2):
            cell = ws.cell(row=r, column=c)
            cell.border = _thin_border()
            if r == HEADER_ROW2 and cell.value:
                cell.font = _header_style()
                cell.fill = _header_fill()
                cell.alignment = Alignment(horizontal='center')

    row_idx = HEADER_ROW2 + 1
    for row in rows:
        values = _session_row_values(row['class_group'].name, row[session_key], groups)
        for c, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=c, value=val)
            cell.border = _thin_border()
        row_idx += 1

    _auto_col_width(ws)


def generate_class_results_excel(data, config=None, sessions=(
        'normale', 'normale_reclamation', 'rattrapage', 'rattrapage_reclamation')):
    """
    `data` : dict retourné par grades.class_results_services.compute_class_results_table()
    — une feuille par session demandée (Session Normale / Session Normale
    après réclamation / Session de Rattrapage / Session de Rattrapage après
    réclamation, les 4 par défaut), chacune à en-têtes fusionnés sur 2
    niveaux, ventilée Homme/Femme, une ligne par classe + une ligne TOTAL.
    `sessions` permet un export par rubrique (une seule session à la fois).
    """
    department = data['department']
    academic_year = data['academic_year']
    semester = data['semester']

    wb = Workbook()
    first = True
    labels = {
        'normale': 'Session Normale', 'normale_reclamation': 'Normale (après récl.)',
        'rattrapage': 'Rattrapage (avant récl.)', 'rattrapage_reclamation': 'Rattrapage (après récl.)',
    }
    titles = {
        'normale': 'Session Normale', 'normale_reclamation': 'Session Normale — après réclamation',
        'rattrapage': 'Session de Rattrapage — avant réclamation',
        'rattrapage_reclamation': 'Session de Rattrapage — après réclamation',
    }
    for session_key in sessions:
        ws = wb.active if first else wb.create_sheet(labels[session_key])
        if first:
            ws.title = labels[session_key]
        first = False
        _write_class_results_sheet(
            ws, session_key, titles[session_key], data['rows'],
            config, department, academic_year, semester,
        )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
