import io
from datetime import date, timedelta
from calendar import monthrange

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.utils.translation import gettext_lazy as _
from django.views.decorators.http import require_http_methods

from .models import StaffPresence, DemandeConge, FichePersonnel
from academic_core.apps.accounts.models import User, Role

# ── Rôles autorisés ──────────────────────────────────────────────────────────
_HR_ROLES = (
    'ADMIN', 'INST_ADMIN', 'SI_ADMIN', 'ASSISTANTE_DG',
    'ADMIN_DIRECTION', 'ADMIN_DE', 'ADMIN_DAF', 'ADMIN_COM', 'ADMIN_RH',
    'ASSISTANTE_DIRECTION', 'ASSISTANTE_DE',
    'CONTROLEUR', 'CIAQ',
    'COMPTABLE', 'TRESORIER_GENERAL', 'CAISSIER',
)

# Rôles pouvant valider/rejeter les congés et modifier les présences
_HR_MANAGERS = (
    'ADMIN', 'INST_ADMIN', 'SI_ADMIN', 'ASSISTANTE_DG',
    'ADMIN_DIRECTION', 'ADMIN_DE', 'ADMIN_DAF', 'ADMIN_COM', 'ADMIN_RH',
    'ASSISTANTE_DIRECTION', 'ASSISTANTE_DE',
    'CONTROLEUR',
)


def _hr_required(view_func):
    """Décorateur : accès réservé aux rôles RH (consultation/gestion globale)."""
    from functools import wraps
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect('accounts:login')
        if request.user.role and request.user.role.name in _HR_ROLES:
            return view_func(request, *args, **kwargs)
        messages.error(request, "Accès refusé : permissions insuffisantes.")
        return redirect('dashboard:index')
    return wrapper


def _staff_required(view_func):
    """
    Décorateur : accès en auto-service, ouvert à tout membre du personnel
    authentifié (pas seulement les rôles RH — cf. décision d'architecture #2 :
    congés, missions, attestations, évaluations… restent consultables/
    actionnables par l'agent concerné lui-même quel que soit son rôle).
    Seuls les étudiants sont exclus ; la logique de gestion (validation,
    création pour un tiers) reste vérifiée à l'intérieur de chaque vue via
    _HR_MANAGERS.
    """
    from functools import wraps
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect('accounts:login')
        if request.user.is_etudiant():
            messages.error(request, "Accès refusé : permissions insuffisantes.")
            return redirect('dashboard:index')
        return view_func(request, *args, **kwargs)
    return wrapper


def _get_staff_queryset(request):
    """
    Retourne le queryset des utilisateurs visibles par le gestionnaire RH.
    Les enseignants VACATAIRES sont exclus : seuls les permanents sont gérés.
    """
    from academic_core.apps.teachers.models import Teacher

    # Exclure les vacataires : utilisateurs ayant un profil Teacher avec statut VACATAIRE
    vacataire_user_ids = Teacher.objects.filter(
        statut=Teacher.STATUT_VACATAIRE
    ).values_list('user_id', flat=True)

    qs = User.objects.select_related('role', 'department', 'direction').filter(
        is_active=True,
    ).exclude(
        role__name='ETUDIANT',
    ).exclude(
        pk__in=vacataire_user_ids,
    ).order_by('last_name', 'first_name')

    # Pas de filtre supplémentaire par faculté : User est routé par tenant, la
    # base courante est déjà celle de l'institut actif (middleware). Filtrer
    # en plus par department__faculty/direction__faculty exclurait à tort tout
    # membre du personnel sans département/direction assigné (ex : un
    # enseignant nouvellement créé sans département).
    return qs


# ── Dashboard RH ─────────────────────────────────────────────────────────────

@login_required
@_hr_required
def hr_dashboard(request):
    today = timezone.now().date()
    year  = int(request.GET.get('year',  today.year))
    month = int(request.GET.get('month', today.month))

    _, nb_jours = monthrange(year, month)
    first_day   = date(year, month, 1)
    last_day    = date(year, month, nb_jours)

    staff_qs = _get_staff_queryset(request)
    nb_staff  = staff_qs.count()

    # Présences du mois
    presences_mois = StaffPresence.objects.filter(
        date__year=year, date__month=month,
        user__in=staff_qs,
    )

    stats_globales = presences_mois.values('statut').annotate(total=Count('id'))
    stats_map = {s['statut']: s['total'] for s in stats_globales}

    # Présences du jour
    presences_jour = StaffPresence.objects.filter(
        date=today, user__in=staff_qs,
    ).select_related('user__role')

    pointes_ids = set(presences_jour.values_list('user_id', flat=True))
    non_pointes = staff_qs.exclude(pk__in=pointes_ids)

    # Stats journalières
    stats_jour = presences_jour.values('statut').annotate(total=Count('id'))
    stats_jour_map = {s['statut']: s['total'] for s in stats_jour}

    # Congés en attente
    conges_en_attente = DemandeConge.objects.filter(
        statut=DemandeConge.STATUT_EN_ATTENTE,
        user__in=staff_qs,
    ).select_related('user').order_by('-created_at')[:10]

    # Navigation mois
    prev_month = month - 1 if month > 1 else 12
    prev_year  = year if month > 1 else year - 1
    next_month = month + 1 if month < 12 else 1
    next_year  = year if month < 12 else year + 1

    MONTHS_FR = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                 'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']

    # Tuiles des nouveaux modules RH — la base courante est déjà celle de
    # l'institut actif (routage tenant), pas de filtre supplémentaire requis.
    from .models import (
        DisciplinaryCase, Evaluation, EvaluationCampaign, Mission, Interim,
        Internship, Onboarding, RecruitmentRequest, Candidate, DocumentRequest,
    )
    discipline_ouverts = DisciplinaryCase.objects.exclude(
        statut__in=[DisciplinaryCase.STATUT_SANCTION, DisciplinaryCase.STATUT_CLASSE]
    ).count()
    campagne_ouverte = EvaluationCampaign.objects.filter(statut=EvaluationCampaign.STATUT_OUVERTE).first()
    evaluations_a_faire = (
        Evaluation.objects.filter(campaign=campagne_ouverte).exclude(statut=Evaluation.STATUT_FINALISEE).count()
        if campagne_ouverte else 0
    )
    missions_en_attente = Mission.objects.filter(statut=Mission.STATUT_SOUMISE).count()
    interims_actifs = Interim.objects.filter(statut=Interim.STATUT_ACTIVE).count()
    stages_en_cours = Internship.objects.filter(statut=Internship.STATUT_EN_COURS).count()
    onboarding_en_cours = Onboarding.objects.filter(statut=Onboarding.STATUT_EN_COURS).count()
    recrutements_ouverts = RecruitmentRequest.objects.exclude(
        statut__in=[RecruitmentRequest.STATUT_POURVUE, RecruitmentRequest.STATUT_REJETEE]
    ).count()
    candidats_actifs = Candidate.objects.exclude(
        statut__in=[Candidate.STATUT_RETENU, Candidate.STATUT_REJETE]
    ).count()
    documents_en_attente = DocumentRequest.objects.exclude(
        statut__in=[DocumentRequest.STATUT_DISPONIBLE, DocumentRequest.STATUT_REJETEE]
    ).count()

    return render(request, 'hr/dashboard.html', {
        'today':             today,
        'year':              year,
        'month':             month,
        'month_label':       MONTHS_FR[month],
        'nb_staff':          nb_staff,
        'nb_jours':          nb_jours,
        'presences_jour':    presences_jour,
        'non_pointes':       non_pointes,
        'stats_jour_map':    stats_jour_map,
        'stats_map':         stats_map,
        'conges_en_attente': conges_en_attente,
        'prev_month': prev_month, 'prev_year': prev_year,
        'next_month': next_month, 'next_year': next_year,
        'STATUT_CHOICES': StaffPresence.STATUT_CHOICES,
        'discipline_ouverts':   discipline_ouverts,
        'campagne_ouverte':     campagne_ouverte,
        'evaluations_a_faire':  evaluations_a_faire,
        'missions_en_attente':  missions_en_attente,
        'interims_actifs':      interims_actifs,
        'stages_en_cours':      stages_en_cours,
        'onboarding_en_cours':  onboarding_en_cours,
        'recrutements_ouverts': recrutements_ouverts,
        'candidats_actifs':     candidats_actifs,
        'documents_en_attente': documents_en_attente,
    })


# ── Pointage journalier (masse) ───────────────────────────────────────────────

@login_required
@_hr_required
def pointage_jour(request):
    """Page pour saisir en masse les présences d'une journée."""
    target_date_str = request.GET.get('date', '')
    try:
        target_date = date.fromisoformat(target_date_str)
    except ValueError:
        target_date = timezone.now().date()

    staff_qs = _get_staff_queryset(request)

    # Présences déjà enregistrées pour ce jour
    existing = {
        p.user_id: p
        for p in StaffPresence.objects.filter(date=target_date, user__in=staff_qs).select_related('user')
    }

    if request.method == 'POST':
        with transaction.atomic():
            for user in staff_qs:
                statut       = request.POST.get(f'statut_{user.pk}', '').strip()
                heure_arr    = request.POST.get(f'arr_{user.pk}', '').strip() or None
                heure_dep    = request.POST.get(f'dep_{user.pk}', '').strip() or None
                justif       = request.POST.get(f'justif_{user.pk}', '').strip()
                doc_file     = request.FILES.get(f'doc_{user.pk}')

                if not statut:
                    continue

                obj, created = StaffPresence.objects.get_or_create(
                    user=user, date=target_date,
                    defaults={'enregistre_par': request.user},
                )
                obj.statut        = statut
                obj.heure_arrivee = heure_arr
                obj.heure_depart  = heure_dep
                obj.justification = justif
                obj.enregistre_par = request.user
                if doc_file:
                    obj.document = doc_file
                obj.save()

        messages.success(request, f"Pointage du {target_date.strftime('%d/%m/%Y')} enregistré avec succès.")
        return redirect(f"{request.path}?date={target_date.isoformat()}")

    # Séparer le personnel en deux groupes pour le template
    staff_pointed   = []  # déjà enregistrés (via scanner ou manuellement)
    staff_unpointed = []  # pas encore pointés
    for emp in staff_qs:
        if emp.pk in existing:
            staff_pointed.append(emp)
        else:
            staff_unpointed.append(emp)

    return render(request, 'hr/pointage_jour.html', {
        'staff_pointed':   staff_pointed,
        'staff_unpointed': staff_unpointed,
        'staff_total':     staff_qs.count(),
        'target_date':     target_date,
        'existing':        existing,
        'STATUT_CHOICES':  StaffPresence.STATUT_CHOICES,
    })


# ── Liste des présences (filtrée) ─────────────────────────────────────────────

@login_required
@_hr_required
def presence_list(request):
    today = timezone.now().date()
    year  = int(request.GET.get('year',  today.year))
    month = int(request.GET.get('month', today.month))
    user_id = request.GET.get('user_id', '')

    staff_qs = _get_staff_queryset(request)

    qs = StaffPresence.objects.filter(
        date__year=year, date__month=month,
        user__in=staff_qs,
    ).select_related('user__role', 'enregistre_par').order_by('-date', 'user__last_name')

    if user_id:
        qs = qs.filter(user_id=user_id)

    _, nb_jours = monthrange(year, month)
    prev_month = month - 1 if month > 1 else 12
    prev_year  = year if month > 1 else year - 1
    next_month = month + 1 if month < 12 else 1
    next_year  = year if month < 12 else year + 1

    MONTHS_FR = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                 'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']

    return render(request, 'hr/presence_list.html', {
        'presences':     qs,
        'staff_list':    staff_qs,
        'year':          year,
        'month':         month,
        'month_label':   MONTHS_FR[month],
        'user_id':       user_id,
        'prev_month': prev_month, 'prev_year': prev_year,
        'next_month': next_month, 'next_year': next_year,
    })


# ── Fiche individuelle d'un employé ──────────────────────────────────────────

@login_required
@_hr_required
def employe_fiche(request, pk):
    staff_qs = _get_staff_queryset(request)
    employe  = get_object_or_404(staff_qs, pk=pk)
    today    = timezone.now().date()
    year     = int(request.GET.get('year',  today.year))
    month    = int(request.GET.get('month', today.month))

    _, nb_jours = monthrange(year, month)
    first_day   = date(year, month, 1)
    last_day    = date(year, month, nb_jours)

    presences = StaffPresence.objects.filter(
        user=employe, date__year=year, date__month=month,
    ).order_by('date')

    stats = presences.values('statut').annotate(total=Count('id'))
    stats_map = {s['statut']: s['total'] for s in stats}

    conges = DemandeConge.objects.filter(user=employe).order_by('-created_at')[:10]

    from .models import Contract, Assignment
    from academic_core.apps.academic_structure.models import Department
    contrats     = Contract.objects.filter(user=employe).prefetch_related('avenants').order_by('-date_debut')
    affectations = Assignment.objects.filter(user=employe).select_related('department').order_by('-date_effet')
    departments  = Department.objects.filter(is_active=True).order_by('name')

    # Calendrier du mois
    cal_days = []
    for day in range(1, nb_jours + 1):
        d = date(year, month, day)
        pres = next((p for p in presences if p.date == d), None)
        cal_days.append({'date': d, 'presence': pres, 'is_weekend': d.weekday() >= 5})

    prev_month = month - 1 if month > 1 else 12
    prev_year  = year if month > 1 else year - 1
    next_month = month + 1 if month < 12 else 1
    next_year  = year if month < 12 else year + 1

    MONTHS_FR = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                 'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']

    return render(request, 'hr/employe_fiche.html', {
        'employe':      employe,
        'presences':    presences,
        'stats_map':    stats_map,
        'conges':       conges,
        'contrats':     contrats,
        'affectations': affectations,
        'departments':  departments,
        'cal_days':    cal_days,
        'year':        year,
        'month':       month,
        'month_label': MONTHS_FR[month],
        'prev_month': prev_month, 'prev_year': prev_year,
        'next_month': next_month, 'next_year': next_year,
        'STATUT_CHOICES': StaffPresence.STATUT_CHOICES,
        'is_manager': request.user.role and request.user.role.name in _HR_MANAGERS,
    })


# ── Modifier / supprimer une présence ─────────────────────────────────────────

@login_required
@_hr_required
def presence_edit(request, pk):
    pres = get_object_or_404(StaffPresence, pk=pk)
    if request.method == 'POST':
        pres.statut        = request.POST.get('statut', pres.statut)
        pres.heure_arrivee = request.POST.get('heure_arrivee') or None
        pres.heure_depart  = request.POST.get('heure_depart') or None
        pres.justification = request.POST.get('justification', '')
        doc = request.FILES.get('document')
        if doc:
            pres.document = doc
        pres.enregistre_par = request.user
        pres.save()
        messages.success(request, "Présence mise à jour.")
        return redirect('hr:employe_fiche', pk=pres.user_id)
    return render(request, 'hr/presence_edit.html', {
        'pres': pres,
        'STATUT_CHOICES': StaffPresence.STATUT_CHOICES,
    })


@login_required
@_hr_required
def presence_delete(request, pk):
    pres = get_object_or_404(StaffPresence, pk=pk)
    user_pk = pres.user_id
    if request.method == 'POST':
        pres.delete()
        messages.success(request, "Pointage supprimé.")
    return redirect('hr:employe_fiche', pk=user_pk)


# ── Congés ────────────────────────────────────────────────────────────────────

@login_required
@_hr_required
def conge_list(request):
    staff_qs = _get_staff_queryset(request)
    statut   = request.GET.get('statut', '')
    user_id  = request.GET.get('user_id', '')

    qs = DemandeConge.objects.filter(user__in=staff_qs).select_related('user', 'valide_par').order_by('-created_at')
    if statut:
        qs = qs.filter(statut=statut)
    if user_id:
        qs = qs.filter(user_id=user_id)

    return render(request, 'hr/conge_list.html', {
        'conges':     qs,
        'staff_list': staff_qs,
        'statut_sel': statut,
        'user_id':    user_id,
        'STATUT_CHOICES': DemandeConge.STATUT_CHOICES,
        'TYPE_CHOICES':   DemandeConge.TYPE_CHOICES,
    })


@login_required
@_hr_required
def conge_create(request):
    staff_qs = _get_staff_queryset(request)
    if request.method == 'POST':
        user_pk    = request.POST.get('user_id')
        type_conge = request.POST.get('type_conge')
        date_debut = request.POST.get('date_debut')
        date_fin   = request.POST.get('date_fin')
        motif      = request.POST.get('motif', '')
        doc        = request.FILES.get('document')

        try:
            emp = User.objects.get(pk=user_pk)
            DemandeConge.objects.create(
                user=emp,
                type_conge=type_conge,
                date_debut=date_debut,
                date_fin=date_fin,
                motif=motif,
                document=doc,
            )
            messages.success(request, f"Demande de congé créée pour {emp.get_full_name()}.")
            return redirect('hr:conge_list')
        except Exception as e:
            messages.error(request, f"Erreur : {e}")

    return render(request, 'hr/conge_form.html', {
        'staff_list':   staff_qs,
        'TYPE_CHOICES': DemandeConge.TYPE_CHOICES,
        'today':        timezone.now().date(),
    })


@login_required
@_hr_required
def conge_review(request, pk):
    """Valider ou rejeter une demande de congé."""
    if not (request.user.role and request.user.role.name in _HR_MANAGERS):
        messages.error(request, "Permission insuffisante pour valider les congés.")
        return redirect('hr:conge_list')

    conge = get_object_or_404(DemandeConge, pk=pk)
    if request.method == 'POST':
        decision         = request.POST.get('decision')
        commentaire      = request.POST.get('commentaire', '')
        conge.commentaire_rh = commentaire
        conge.valide_par     = request.user
        conge.valide_le      = timezone.now()

        if decision == 'APPROUVE':
            conge.statut = DemandeConge.STATUT_APPROUVE
            # Créer automatiquement les StaffPresence pour la période
            _auto_create_conge_presences(conge, request.user)
            messages.success(request, "Demande approuvée. Les présences ont été générées automatiquement.")
        elif decision == 'REJETE':
            conge.statut = DemandeConge.STATUT_REJETE
            messages.warning(request, "Demande rejetée.")

        conge.save()
        return redirect('hr:conge_list')

    return render(request, 'hr/conge_review.html', {'conge': conge})


def _auto_create_conge_presences(conge, enregistre_par):
    """Crée les StaffPresence CONGE pour toute la période approuvée."""
    current = conge.date_debut
    while current <= conge.date_fin:
        if current.weekday() < 5:  # jours ouvrables uniquement
            StaffPresence.objects.get_or_create(
                user=conge.user,
                date=current,
                defaults={
                    'statut':        StaffPresence.STATUT_CONGE,
                    'justification': f"Congé approuvé — {conge.get_type_conge_display()}",
                    'enregistre_par': enregistre_par,
                },
            )
        current += timedelta(days=1)


# ── Export Excel mensuel ──────────────────────────────────────────────────────

@login_required
@_hr_required
def export_excel(request):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, "La bibliothèque openpyxl est requise pour l'export Excel.")
        return redirect('hr:presence_list')

    today = timezone.now().date()
    year  = int(request.GET.get('year',  today.year))
    month = int(request.GET.get('month', today.month))
    _, nb_jours = monthrange(year, month)

    MONTHS_FR = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                 'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']

    staff_qs = _get_staff_queryset(request)
    presences = StaffPresence.objects.filter(
        date__year=year, date__month=month, user__in=staff_qs,
    ).select_related('user')
    pres_map = {}
    for p in presences:
        pres_map.setdefault(p.user_id, {})[p.date.day] = p

    wb  = Workbook()
    ws  = wb.active
    ws.title = f"Présences {MONTHS_FR[month]} {year}"

    NAVY  = PatternFill('solid', fgColor='0D2244')
    GREEN = PatternFill('solid', fgColor='16A34A')
    RED   = PatternFill('solid', fgColor='DC2626')
    AMBER = PatternFill('solid', fgColor='D97706')
    BLUE  = PatternFill('solid', fgColor='0EA5E9')
    GRAY  = PatternFill('solid', fgColor='E2E8F0')

    STATUT_FILL = {
        'PRESENT': GREEN, 'ABSENT': RED, 'RETARD': AMBER,
        'CONGE': BLUE, 'DEMI_JOURNEE': GRAY,
        'TELETRAVAIL': PatternFill('solid', fgColor='7C3AED'),
        'FERIE': PatternFill('solid', fgColor='64748B'),
    }
    STATUT_CODE = {
        'PRESENT': 'P', 'ABSENT': 'A', 'RETARD': 'R',
        'CONGE': 'C', 'DEMI_JOURNEE': 'D', 'TELETRAVAIL': 'T', 'FERIE': 'F',
    }

    white_bold  = Font(color='FFFFFF', bold=True, size=10)
    dark_bold   = Font(bold=True, size=10)
    center_al   = Alignment(horizontal='center', vertical='center')
    thin        = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    # Titre
    ws.merge_cells('A1:C1')
    ws['A1'] = f"Registre des Présences — {MONTHS_FR[month]} {year}"
    ws['A1'].font      = Font(bold=True, size=13, color='FFFFFF')
    ws['A1'].fill      = NAVY
    ws['A1'].alignment = center_al
    ws.row_dimensions[1].height = 28

    # En-têtes
    row = 3
    ws.cell(row, 1, 'N°').font = white_bold
    ws.cell(row, 2, 'Employé').font = white_bold
    ws.cell(row, 3, 'Rôle').font = white_bold
    for c in [1, 2, 3]:
        ws.cell(row, c).fill = NAVY
        ws.cell(row, c).alignment = center_al
        ws.cell(row, c).border = thin

    for day in range(1, nb_jours + 1):
        col = day + 3
        cell = ws.cell(row, col, str(day))
        d = date(year, month, day)
        cell.fill      = GRAY if d.weekday() >= 5 else NAVY
        cell.font      = white_bold
        cell.alignment = center_al
        cell.border    = thin

    # Colonnes totaux
    tot_cols = {'P': nb_jours + 4, 'A': nb_jours + 5, 'R': nb_jours + 6, 'C': nb_jours + 7}
    labels = {'P': 'Présent', 'A': 'Absent', 'R': 'Retard', 'C': 'Congé'}
    for code, col in tot_cols.items():
        cell = ws.cell(row, col, labels[code])
        cell.fill = NAVY; cell.font = white_bold
        cell.alignment = center_al; cell.border = thin

    # Données
    for idx, emp in enumerate(staff_qs, 1):
        row += 1
        bg = PatternFill('solid', fgColor='F8FAFC') if idx % 2 == 0 else PatternFill('solid', fgColor='FFFFFF')
        ws.cell(row, 1, idx).alignment = center_al
        ws.cell(row, 2, emp.get_full_name())
        ws.cell(row, 3, emp.role.get_name_display() if emp.role else '')
        for c in [1, 2, 3]:
            ws.cell(row, c).fill   = bg
            ws.cell(row, c).border = thin

        totaux = {'P': 0, 'A': 0, 'R': 0, 'C': 0}
        for day in range(1, nb_jours + 1):
            col  = day + 3
            pres = pres_map.get(emp.pk, {}).get(day)
            cell = ws.cell(row, col)
            cell.border    = thin
            cell.alignment = center_al
            if pres:
                code = STATUT_CODE.get(pres.statut, '?')
                cell.value = code
                cell.fill  = STATUT_FILL.get(pres.statut, bg)
                cell.font  = Font(bold=True, size=9,
                                  color='FFFFFF' if pres.statut not in ('RETARD',) else '000000')
                if code in totaux:
                    totaux[code] += 1
            else:
                cell.fill = bg

        for code, col in tot_cols.items():
            c = ws.cell(row, col, totaux[code])
            c.alignment = center_al
            c.border    = thin
            c.font      = dark_bold

    # Légende
    row += 2
    ws.cell(row, 1, 'Légende :').font = dark_bold
    legend = [('P', 'Présent', '16A34A'), ('A', 'Absent', 'DC2626'),
              ('R', 'Retard', 'D97706'), ('C', 'Congé', '0EA5E9'),
              ('D', 'Demi-journée', 'E2E8F0'), ('T', 'Télétravail', '7C3AED'), ('F', 'Férié', '64748B')]
    for i, (code, label, color) in enumerate(legend):
        col = i + 2
        c   = ws.cell(row, col, f"{code} = {label}")
        c.fill   = PatternFill('solid', fgColor=color)
        c.font   = Font(size=8, bold=True, color='FFFFFF' if color not in ('E2E8F0',) else '000000')
        c.border = thin
        c.alignment = center_al

    # Largeurs colonnes
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['C'].width = 20
    for day in range(1, nb_jours + 1):
        ws.column_dimensions[get_column_letter(day + 3)].width = 4
    for col in tot_cols.values():
        ws.column_dimensions[get_column_letter(col)].width = 10

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"presences_{year}_{month:02d}.xlsx"
    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ── AJAX : toggle statut rapide ────────────────────────────────────────────────

@login_required
@_hr_required
@require_http_methods(["POST"])
def ajax_set_presence(request):
    """Met à jour ou crée un pointage via AJAX (utilisé dans le tableau de bord)."""
    import json
    try:
        body    = json.loads(request.body)
        user_pk = int(body['user_id'])
        d       = date.fromisoformat(body['date'])
        statut  = body['statut']
    except (KeyError, ValueError, TypeError):
        return JsonResponse({'error': 'Paramètres invalides'}, status=400)

    staff_qs = _get_staff_queryset(request)
    emp = get_object_or_404(staff_qs, pk=user_pk)

    if statut == 'DELETE':
        StaffPresence.objects.filter(user=emp, date=d).delete()
        return JsonResponse({'ok': True, 'deleted': True})

    if statut not in dict(StaffPresence.STATUT_CHOICES):
        return JsonResponse({'error': 'Statut invalide'}, status=400)

    heure_arr = body.get('heure_arrivee', None)
    heure_dep = body.get('heure_depart', None)
    justif    = body.get('justification', None)

    defaults = {'statut': statut, 'enregistre_par': request.user}
    if heure_arr is not None:
        defaults['heure_arrivee'] = heure_arr or None
    if heure_dep is not None:
        defaults['heure_depart'] = heure_dep or None
    if justif is not None:
        defaults['justification'] = justif

    obj, _ = StaffPresence.objects.update_or_create(
        user=emp, date=d,
        defaults=defaults,
    )
    return JsonResponse({
        'ok': True,
        'statut': obj.statut,
        'statut_label': obj.get_statut_display(),
        'color': obj.statut_color,
    })


# ── Génération de matricule employé ──────────────────────────────────────────

def _generate_matricule(user, sigle='ISI'):
    """
    Génère un matricule unique de type : 42Emp-24-39/ISI
    Format : {pk}Emp-{year_2digits}-{sequence}/{sigle}
    """
    year_2 = str(timezone.now().year)[-2:]
    # Séquence = nombre d'utilisateurs avec matricule déjà créé + 1
    seq = User.objects.filter(
        matricule_employe__isnull=False
    ).exclude(matricule_employe='').count() + 1
    return f"{user.pk}Emp-{year_2}-{seq:02d}/{sigle}"


# ── Cartes du personnel ───────────────────────────────────────────────────────

@login_required
@_hr_required
def cartes_personnel(request):
    """Liste du personnel avec statut de carte et bouton d'impression."""
    staff_qs = _get_staff_queryset(request)

    # Récupérer le sigle de l'institut courant
    institut = getattr(request, 'active_faculty', None) or getattr(request, 'active_institute', None)
    sigle = ''
    if institut:
        try:
            from academic_core.apps.academic_structure.models import InstitutConfig
            cfg = InstitutConfig.objects.filter(faculty=institut).first()
            sigle = cfg.sigle if cfg else ''
        except Exception:
            pass

    # Auto-générer les matricules manquants
    if request.method == 'POST' and request.POST.get('action') == 'gen_all':
        for emp in staff_qs.filter(
            Q(matricule_employe__isnull=True) | Q(matricule_employe='')
        ):
            emp.matricule_employe = _generate_matricule(emp, sigle or 'ISI')
            emp.save(update_fields=['matricule_employe'])
        messages.success(request, "Matricules générés pour tous les employés sans carte.")
        return redirect('hr:cartes_personnel')

    return render(request, 'hr/cartes_personnel.html', {
        'staff_list': staff_qs,
        'sigle': sigle,
    })


@login_required
@_hr_required
def carte_employe(request, pk):
    """Vue d'impression de la carte pour un employé (recto + verso)."""
    staff_qs = _get_staff_queryset(request)
    emp      = get_object_or_404(staff_qs, pk=pk)

    # Auto-générer le matricule si absent
    if not emp.matricule_employe:
        institut = getattr(request, 'active_faculty', None)
        sigle = 'ISI'
        if institut:
            try:
                from academic_core.apps.academic_structure.models import InstitutConfig
                cfg = InstitutConfig.objects.filter(faculty=institut).first()
                sigle = cfg.sigle if cfg else 'ISI'
            except Exception:
                pass
        emp.matricule_employe = _generate_matricule(emp, sigle)
        emp.save(update_fields=['matricule_employe'])

    # Récupérer la config de l'institut pour la carte
    institut_cfg = None
    try:
        from academic_core.apps.academic_structure.models import InstitutConfig
        if hasattr(request, 'active_faculty') and request.active_faculty:
            institut_cfg = InstitutConfig.objects.filter(
                faculty=request.active_faculty
            ).first()
    except Exception:
        pass

    # Jeton de pointage + fiche personnel (créés à la volée si absents)
    token = emp.get_or_create_staff_qr_token()
    fiche, _created = FichePersonnel.objects.get_or_create(user=emp)

    # Générer le QR code en base64 (fonctionne sans CDN externe)
    # Le QR encode l'URL de pointage self-service (et non plus le matricule en clair)
    checkin_url = request.build_absolute_uri(
        f"/hr/pointer/{token}/"
    )
    qr_b64 = None
    try:
        import qrcode, io, base64
        qr = qrcode.QRCode(
            version=None,
            error_correction=qrcode.constants.ERROR_CORRECT_H,
            box_size=6,
            border=2,
        )
        qr.add_data(checkin_url)
        qr.make(fit=True)
        img = qr.make_image(fill_color='#0d2244', back_color='#f8fafc')
        buf = io.BytesIO()
        img.save(buf, format='PNG')
        qr_b64 = 'data:image/png;base64,' + base64.b64encode(buf.getvalue()).decode()
    except Exception:
        pass

    return render(request, 'hr/carte_employe.html', {
        'emp': emp,
        'fiche': fiche,
        'institut_cfg': institut_cfg,
        'print_mode': request.GET.get('print') == '1',
        'qr_b64': qr_b64,
        'checkin_url': checkin_url,
    })


@login_required
def ma_carte_personnel(request):
    """
    Permet à chaque membre du personnel de consulter et imprimer sa propre carte.
    Accessible à tous les utilisateurs authentifiés non-étudiants.
    """
    user = request.user
    if user.is_etudiant():
        messages.error(request, "Cette page n'est pas disponible pour les étudiants.")
        return redirect('dashboard:index')

    # Le Super Admin n'est le personnel d'aucun institut en particulier — sans
    # ce garde-fou, FichePersonnel.objects.get_or_create() ci-dessous tenterait
    # d'insérer sa carte dans la base de l'institut actuellement actif (s'il y
    # a « Accédé »), où sa ligne User n'existe jamais localement (son compte ne
    # vit que dans 'default'), ce qui lève IntegrityError: FOREIGN KEY
    # constraint failed.
    if user.is_super_admin():
        messages.info(request, "Cette page n'est pas disponible pour le Super Administrateur.")
        return redirect('dashboard:index')

    # Auto-générer le matricule si absent
    if not user.matricule_employe:
        institut = getattr(request, 'active_faculty', None)
        sigle = 'ISI'
        if institut:
            try:
                from academic_core.apps.academic_structure.models import InstitutConfig
                cfg = InstitutConfig.objects.filter(faculty=institut).first()
                sigle = cfg.sigle if cfg else 'ISI'
            except Exception:
                pass
        user.matricule_employe = _generate_matricule(user, sigle)
        user.save(update_fields=['matricule_employe'])

    # Config institut
    institut_cfg = None
    try:
        from academic_core.apps.academic_structure.models import InstitutConfig
        if hasattr(request, 'active_faculty') and request.active_faculty:
            institut_cfg = InstitutConfig.objects.filter(faculty=request.active_faculty).first()
    except Exception:
        pass

    # Filet de sécurité : si la copie locale de l'utilisateur est absente de
    # la base tenant actuellement active (ex : synchronisation initiale
    # échouée silencieusement), la réparer avant l'insertion FichePersonnel
    # ci-dessous, plutôt que de planter sur une IntegrityError FK.
    from academic_core.db_router import get_current_db
    current_alias = get_current_db()
    if current_alias and current_alias != 'default' and not User.objects.filter(pk=user.pk).exists():
        from academic_core.apps.accounts.db_utils import sync_user_to_institute_db
        sync_user_to_institute_db(user, current_alias)

    # Jeton QR + fiche
    token = user.get_or_create_staff_qr_token()
    fiche, _ = FichePersonnel.objects.get_or_create(user=user)

    # QR Code
    checkin_url = request.build_absolute_uri(f"/hr/pointer/{token}/")
    qr_b64 = None
    try:
        import qrcode, base64
        qr = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_H,
                            box_size=6, border=2)
        qr.add_data(checkin_url)
        qr.make(fit=True)
        img = qr.make_image(fill_color='#0d2244', back_color='#f8fafc')
        buf = io.BytesIO()
        img.save(buf, format='PNG')
        qr_b64 = 'data:image/png;base64,' + base64.b64encode(buf.getvalue()).decode()
    except Exception:
        pass

    return render(request, 'hr/carte_employe.html', {
        'emp':          user,
        'fiche':        fiche,
        'institut_cfg': institut_cfg,
        'print_mode':   request.GET.get('print') == '1',
        'qr_b64':       qr_b64,
        'checkin_url':  checkin_url,
        'self_view':    True,
    })


@login_required
def ma_carte_pointage(request):
    """
    Page 'Pointage à l'accueil' — commune aux étudiants et au personnel : la
    personne active sa propre caméra pour scanner le QR Code affiché à
    l'accueil (accueil_qr_display) et enregistrer son passage, sans passer
    par un agent de contrôle. Étudiant -> statut de paiement affiché en
    carte ; personnel -> arrivée/départ (voir accueil_self_checkin, qui
    branche sur le rôle de la même façon).
    """
    user = request.user

    # Voir le commentaire dans ma_carte_personnel : le Super Admin n'est le
    # personnel d'aucun institut en particulier — le pointage self-scan
    # échouerait de toute façon (StaffPresence n'a pas de copie locale de son
    # compte dans la base tenant active), autant ne pas l'y amener.
    if user.is_super_admin():
        messages.info(request, "Cette page n'est pas disponible pour le Super Administrateur.")
        return redirect('dashboard:index')

    return render(request, 'hr/ma_carte_pointage.html', {'employe': user, 'is_student': user.is_etudiant()})


# ─────────────────────────────────────────────────────────────────────────────
# CONTRÔLE ACCUEIL — page de scan du personnel
# ─────────────────────────────────────────────────────────────────────────────

_ACCUEIL_PAGE_ROLES = (
    'CONTROLE_ACCUEIL', 'ADMIN', 'INST_ADMIN', 'ADMIN_DIRECTION',
    'ADMIN_DE', 'ADMIN_DAF', 'ADMIN_COM', 'ADMIN_RH',
    'ASSISTANTE_DE',
)


@login_required
def controle_accueil_personnel(request):
    """Page de scan QR dédiée au pointage du personnel — rôle CONTROLE_ACCUEIL."""
    role_name = request.user.role.name if request.user.role else ''
    if role_name not in _ACCUEIL_PAGE_ROLES:
        messages.error(request, "Accès refusé.")
        return redirect('dashboard:index')
    return render(request, 'hr/controle_accueil_personnel.html')


@login_required
def accueil_qr_display(request):
    """
    Affiche en plein écran le QR Code du contrôle d'accueil : chaque membre du
    personnel le scanne avec sa propre caméra depuis 'Ma carte de pointage'
    pour enregistrer son arrivée/départ (voir checkin_views.accueil_self_checkin).
    Jeton valable pour la journée en cours (checkin_views.daily_accueil_token),
    régénéré automatiquement chaque jour — pas de rafraîchissement manuel requis.
    """
    role_name = request.user.role.name if request.user.role else ''
    if role_name not in _ACCUEIL_PAGE_ROLES:
        messages.error(request, "Accès refusé.")
        return redirect('dashboard:index')

    import qrcode, base64
    from .checkin_views import daily_accueil_token
    from academic_core.apps.attendance.views import _get_accessible_base_url

    token       = daily_accueil_token()
    base_url    = _get_accessible_base_url(request)
    checkin_url = f"{base_url.rstrip('/')}/hr/pointer-auth/{token}/"

    qr = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_H,
                        box_size=8, border=2)
    qr.add_data(checkin_url)
    qr.make(fit=True)
    img = qr.make_image(fill_color='#0d2244', back_color='#ffffff')
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    qr_b64 = base64.b64encode(buf.getvalue()).decode()

    return render(request, 'hr/accueil_qr_display.html', {
        'qr_b64':      qr_b64,
        'checkin_url': checkin_url,
        'today':       timezone.localdate(),
    })


# ─────────────────────────────────────────────────────────────────────────────
# RAPPORT MENSUEL DE POINTAGE
# ─────────────────────────────────────────────────────────────────────────────

@login_required
@_hr_required
def rapport_mensuel_pointage(request):
    """
    Rapport mensuel de pointage du personnel.
    Filtre par mois + année. Pour chaque employé : jours ouvrés,
    présences, retards, absences.
    """
    from calendar import monthrange
    today = timezone.localdate()

    try:
        mois  = int(request.GET.get('mois',  today.month))
        annee = int(request.GET.get('annee', today.year))
    except (ValueError, TypeError):
        mois, annee = today.month, today.year

    mois  = max(1, min(12, mois))
    annee = max(2020, min(today.year + 1, annee))

    # Jours ouvrés du mois (lundi–vendredi)
    nb_days = monthrange(annee, mois)[1]
    jours_ouvres = sum(
        1 for d in range(1, nb_days + 1)
        if date(annee, mois, d).weekday() < 5
    )

    # Présences du mois
    presences_qs = StaffPresence.objects.filter(
        date__year=annee,
        date__month=mois,
    ).select_related('user', 'user__role')

    # Regrouper par employé
    from collections import defaultdict
    par_employe = defaultdict(lambda: {'present': 0, 'retard': 0, 'absent': 0, 'autre': 0})
    for p in presences_qs:
        if p.statut == StaffPresence.STATUT_PRESENT:
            par_employe[p.user]['present'] += 1
        elif p.statut == StaffPresence.STATUT_RETARD:
            par_employe[p.user]['retard'] += 1
        elif p.statut == StaffPresence.STATUT_ABSENT:
            par_employe[p.user]['absent'] += 1
        else:
            par_employe[p.user]['autre'] += 1

    # Récupérer tous les employés actifs (sans étudiants)
    staff_qs = _get_staff_queryset(request)
    rows = []
    for emp in staff_qs:
        stats = par_employe.get(emp, {'present': 0, 'retard': 0, 'absent': 0, 'autre': 0})
        total_enregistre = stats['present'] + stats['retard'] + stats['absent'] + stats['autre']
        rows.append({
            'emp':            emp,
            'present':        stats['present'],
            'retard':         stats['retard'],
            'absent':         stats['absent'],
            'autre':          stats['autre'],
            'total':          total_enregistre,
            'jours_ouvres':   jours_ouvres,
            'non_enregistre': max(0, jours_ouvres - total_enregistre),
        })

    MOIS_FR = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
               'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']

    return render(request, 'hr/rapport_mensuel_pointage.html', {
        'rows':          rows,
        'mois':          mois,
        'annee':         annee,
        'mois_label':    MOIS_FR[mois],
        'jours_ouvres':  jours_ouvres,
        'annees':        range(2020, today.year + 2),
        'mois_choices':  [(i, MOIS_FR[i]) for i in range(1, 13)],
    })


@login_required
@_hr_required
@require_http_methods(["POST"])
def ajax_gen_matricule(request):
    """Génère ou regénère le matricule d'un employé."""
    import json as _json
    try:
        body = _json.loads(request.body)
        user_pk = int(body['user_id'])
    except (KeyError, ValueError, TypeError):
        return JsonResponse({'error': 'Paramètres invalides'}, status=400)

    staff_qs = _get_staff_queryset(request)
    emp = get_object_or_404(staff_qs, pk=user_pk)
    sigle = body.get('sigle', 'ISI')
    emp.matricule_employe = _generate_matricule(emp, sigle)
    emp.save(update_fields=['matricule_employe'])
    return JsonResponse({'ok': True, 'matricule': emp.matricule_employe})
