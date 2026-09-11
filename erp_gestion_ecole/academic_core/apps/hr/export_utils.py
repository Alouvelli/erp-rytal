"""
Générateurs génériques d'export PDF / Word pour les listes du module Ressources Humaines.
"""
import io

from django.http import HttpResponse
from django.utils import timezone


def build_pdf_table_response(title, subtitle, headers, rows, filename,
                              orientation='portrait', institut_config=None,
                              col_widths=None):
    """Génère un PDF tabulaire générique (portrait ou paysage)."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.lib.enums import TA_CENTER
    from academic_core.pdf_utils import logo_image, watermark_canvas

    navy = colors.HexColor('#00173B')
    gold = colors.HexColor('#D4AF37')
    grey = colors.HexColor('#64748B')

    pagesize = landscape(A4) if orientation == 'landscape' else A4
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=pagesize,
        rightMargin=1.2*cm, leftMargin=1.2*cm, topMargin=1*cm, bottomMargin=1*cm,
    )

    ss = getSampleStyleSheet()
    sc = ParagraphStyle
    s_inst = sc('inst', parent=ss['Normal'], fontSize=11, fontName='Helvetica-Bold', textColor=navy)
    s_ttl  = sc('ttl',  parent=ss['Normal'], alignment=TA_CENTER, fontSize=13, fontName='Helvetica-Bold', textColor=navy)
    s_sub  = sc('sub',  parent=ss['Normal'], alignment=TA_CENTER, fontSize=9, textColor=grey)
    s_foot = sc('foot', parent=ss['Normal'], alignment=TA_CENTER, fontSize=7, textColor=grey)
    s_wm   = sc('wm',   parent=ss['Normal'], alignment=TA_CENTER, fontSize=6, textColor=colors.HexColor('#94A3B8'))
    s_cell = sc('cell', parent=ss['Normal'], fontSize=7.5)
    s_hcell = sc('hcell', parent=ss['Normal'], fontSize=7.5, fontName='Helvetica-Bold', textColor=colors.white)

    nom_inst = getattr(institut_config, 'nom', None) or "Institut Supérieur d'Informatique"
    logo = logo_image(config=institut_config, width=1.6*cm, height=1.6*cm)
    hdr_cells = [logo or '', Paragraph(f"<b>{nom_inst}</b>", s_inst)]
    hdr_tbl = Table([hdr_cells], colWidths=[2*cm, None])
    hdr_tbl.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))

    story = [
        hdr_tbl,
        Spacer(1, .2*cm),
        HRFlowable(width='100%', thickness=2, color=navy),
        HRFlowable(width='100%', thickness=1, color=gold, spaceAfter=8),
        Paragraph(title, s_ttl),
    ]
    if subtitle:
        story.append(Paragraph(subtitle, s_sub))
    story.append(Spacer(1, .4*cm))

    table_data = [[Paragraph(h, s_hcell) for h in headers]]
    for row in rows:
        table_data.append([Paragraph(str(v) if v not in (None, '') else '—', s_cell) for v in row])

    tbl = Table(table_data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), navy),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#E2E8F0')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(tbl)
    story.append(Spacer(1, .5*cm))
    story.append(HRFlowable(width='100%', thickness=0.5, color=colors.lightgrey))
    story.append(Paragraph(
        f"{len(rows)} ligne(s) — Document généré le {timezone.now().strftime('%d/%m/%Y à %H:%M')} — {nom_inst}",
        s_foot,
    ))
    doc.build(story, onFirstPage=watermark_canvas, onLaterPages=watermark_canvas)
    buffer.seek(0)
    resp = HttpResponse(buffer.read(), content_type='application/pdf')
    resp['Content-Disposition'] = f'attachment; filename="{filename}.pdf"'
    return resp


def build_docx_table_response(title, subtitle, headers, rows, filename,
                               orientation='portrait', institut_config=None):
    """Génère un document Word (.docx) tabulaire générique (portrait ou paysage)."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT

    doc = Document()
    sec = doc.sections[0]
    if orientation == 'landscape':
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width

    nom_inst = getattr(institut_config, 'nom', None) or "Institut Supérieur d'Informatique"

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run(nom_inst)
    r_inst.bold = True
    r_inst.font.size = Pt(13)
    r_inst.font.color.rgb = RGBColor(0x00, 0x17, 0x3B)

    p_ttl = doc.add_paragraph()
    p_ttl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ttl = p_ttl.add_run(title)
    r_ttl.bold = True
    r_ttl.font.size = Pt(15)

    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_sub = p_sub.add_run(subtitle)
        r_sub.font.size = Pt(10)
        r_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    for i, htext in enumerate(headers):
        hdr_cells[i].text = str(htext)
        for run in hdr_cells[i].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)

    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val) if val not in (None, '') else '—'
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(8.5)

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = footer.add_run(
        f"{len(rows)} ligne(s) — Document généré le {timezone.now().strftime('%d/%m/%Y à %H:%M')} — {nom_inst}"
    )
    r_foot.font.size = Pt(7.5)
    r_foot.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    from academic_core.pdf_utils import add_docx_watermark
    add_docx_watermark(doc)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    resp['Content-Disposition'] = f'attachment; filename="{filename}.docx"'
    return resp
