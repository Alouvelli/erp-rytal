"""
Planning de l'emploi du temps — grille hebdomadaire par classe, enseignant ou
salle (affichage écran + export PDF/Excel).

La grille suppose un axe unique sélectionné (une classe, OU un enseignant, OU
une salle) : sur cet axe, deux créneaux ne peuvent jamais se chevaucher (les
conflits sont bloqués à la saisie — voir TimetableEntry.check_conflicts), ce
qui garantit au plus une entrée par (jour, heure).
"""
import html
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect

from .models import TimetableEntry


def _esc(s):
    return html.escape(str(s)) if s else ''


def build_planning_grid(entries):
    """
    Construit une grille hebdomadaire à partir d'une liste de TimetableEntry
    (déjà filtrée sur un axe unique — classe, enseignant ou salle).

    Retourne None si `entries` est vide, sinon :
    {
        'days': [{'num': 1, 'label': 'Lundi'}, ...],
        'rows': [{'start': time, 'end': time, 'cells': [cell, ...]}, ...],
    }
    où chaque `cell` (aligné sur `days`) vaut :
      - None                          → couvert par le rowspan d'une ligne au-dessus (ne rien afficher)
      - {'entry': None, 'span': 1}     → créneau libre
      - {'entry': <TimetableEntry>, 'span': n} → séance occupant n lignes
    """
    entries = list(entries)
    if not entries:
        return None

    boundaries = sorted({e.start_time for e in entries} | {e.end_time for e in entries})
    if len(boundaries) < 2:
        return None
    b_index = {t: i for i, t in enumerate(boundaries)}

    days = [1, 2, 3, 4, 5, 6]
    if any(e.day_of_week == 7 for e in entries):
        days.append(7)
    # str() : DAY_CHOICES utilise gettext_lazy, dont les proxies de traduction
    # ne sont pas acceptés tels quels par openpyxl (Cannot convert ... to Excel).
    day_labels = {k: str(v) for k, v in TimetableEntry.DAY_CHOICES}

    # (jour, ligne) -> {'entry':e,'span':n} sur la ligne de départ, 'SKIP' sur les lignes couvertes
    cell_map = {}
    for e in entries:
        start_idx = b_index[e.start_time]
        end_idx   = b_index[e.end_time]
        span = max(1, end_idx - start_idx)
        cell_map[(e.day_of_week, start_idx)] = {'entry': e, 'span': span}
        for r in range(start_idx + 1, end_idx):
            cell_map[(e.day_of_week, r)] = 'SKIP'

    rows = []
    for row_idx in range(len(boundaries) - 1):
        row_cells = []
        for d in days:
            info = cell_map.get((d, row_idx))
            if info == 'SKIP':
                row_cells.append(None)
            elif info:
                row_cells.append(info)
            else:
                row_cells.append({'entry': None, 'span': 1})
        rows.append({
            'start': boundaries[row_idx],
            'end':   boundaries[row_idx + 1],
            'cells': row_cells,
        })

    return {
        'days': [{'num': d, 'label': day_labels.get(d, '')} for d in days],
        'rows': rows,
    }


def _filtered_entries(request):
    """Reproduit les filtres de TimetableIndexView, restreints à un axe unique."""
    from academic_core.apps.academic_structure.models import Semester

    qs = TimetableEntry.objects.select_related(
        'class_group', 'subject', 'teacher__user', 'room', 'semester'
    ).filter(is_active=True)

    user = request.user
    if user.is_etudiant():
        sp = getattr(user, 'student_profile', None)
        if sp:
            current = sp.current_enrollment()
            if current:
                qs = qs.filter(class_group=current.class_group)
    elif user.is_enseignant():
        qs = qs.filter(teacher__user=user)

    semester_id = request.GET.get('semester')
    class_id    = request.GET.get('class_group')
    teacher_id  = request.GET.get('teacher')
    room_id     = request.GET.get('room')

    # Même comportement que TimetableIndexView.get_queryset() : ne filtrer par
    # semestre que si explicitement sélectionné, pour rester cohérent avec ce
    # que l'utilisateur voit à l'écran (calendrier/liste/grille).
    semester = None
    if semester_id:
        semester = Semester.objects.filter(pk=semester_id).select_related('academic_year').first()
        if semester:
            qs = qs.filter(semester=semester)

    # Déterminer l'axe réellement actif (une classe, un enseignant OU une salle)
    from academic_core.apps.academic_structure.models import Class as ClassGroup
    from academic_core.apps.teachers.models import Teacher
    from academic_core.apps.rooms.models import Room

    axis_label, axis_obj = None, None
    if class_id and not user.is_etudiant():
        qs = qs.filter(class_group_id=class_id)
        axis_obj = ClassGroup.objects.filter(pk=class_id).first()
        axis_label = 'Classe'
    elif teacher_id and user.can_manage_dept():
        qs = qs.filter(teacher_id=teacher_id)
        axis_obj = Teacher.objects.select_related('user').filter(pk=teacher_id).first()
        axis_label = 'Enseignant'
    elif room_id and user.can_manage_dept():
        qs = qs.filter(room_id=room_id)
        axis_obj = Room.objects.filter(pk=room_id).first()
        axis_label = 'Salle'
    elif user.is_enseignant():
        axis_obj = getattr(user, 'teacher_profile', None)
        axis_label = 'Enseignant'
    elif user.is_etudiant():
        sp = getattr(user, 'student_profile', None)
        current = sp.current_enrollment() if sp else None
        axis_obj = current.class_group if current else None
        axis_label = 'Classe'

    return qs.order_by('day_of_week', 'start_time'), semester, axis_label, axis_obj


def generate_planning_pdf_bytes(grid, axis_label, axis_obj, semester, config=None):
    """
    Génère le PDF (A4 paysage) de la grille hebdomadaire déjà construite
    (`build_planning_grid`) pour un axe (classe/enseignant/salle) — factorisé
    hors de `planning_pdf` pour être réutilisable sans objet `request` (ex.
    envoi par email du planning d'une classe à ses étudiants, ou du planning
    personnel d'un enseignant — voir send_class_planning_emails /
    send_teacher_planning_emails). Retourne un buffer BytesIO positionné au
    début.
    """
    import io
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from academic_core.pdf_utils import watermark_canvas, logo_image

    inst_name = (config.nom if config and getattr(config, 'nom', None) else '') or ''

    navy   = colors.HexColor('#1e3a5f')
    border = colors.HexColor('#e2e8f0')
    lgrey  = colors.HexColor('#f8fafc')

    buf = io.BytesIO()
    pagesize = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=pagesize,
                             leftMargin=1*cm, rightMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm)
    usable = pagesize[0] - 2*cm

    s_title = ParagraphStyle('T', fontSize=14, fontName='Helvetica-Bold', textColor=navy, alignment=TA_CENTER)
    s_sub   = ParagraphStyle('S', fontSize=9, textColor=colors.HexColor('#64748b'), alignment=TA_CENTER)
    s_cell  = ParagraphStyle('C', fontSize=7.5, leading=9.5, textColor=navy, fontName='Helvetica-Bold', alignment=TA_CENTER)
    s_time  = ParagraphStyle('TM', fontSize=7.5, fontName='Helvetica-Bold', textColor=colors.HexColor('#64748b'), alignment=TA_CENTER)
    s_hdr   = ParagraphStyle('H', fontSize=8.5, fontName='Helvetica-Bold', textColor=colors.white, alignment=TA_CENTER)

    _logo = logo_image(width=1.6*cm, height=1.6*cm, config=config)
    logo_row = Table([[
        _logo or '',
        Paragraph((inst_name or '').upper(), ParagraphStyle(
            'InstName', fontSize=10, fontName='Helvetica-Bold', textColor=navy, alignment=TA_CENTER,
        )),
    ]], colWidths=[2.2*cm, None])
    logo_row.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))

    if axis_label == 'Classe' and axis_obj is not None:
        filiere_name = axis_obj.program.name if getattr(axis_obj, 'program', None) else ''
        axis_title = f"{filiere_name} - {axis_obj.name}" if filiere_name else axis_obj.name
    else:
        axis_title = f"{axis_label.upper()} : {axis_obj}"

    story = [
        logo_row,
        HRFlowable(width='100%', thickness=1, color=navy, spaceAfter=4),
        Paragraph(f"EMPLOI DU TEMPS — {axis_title}", s_title),
        Paragraph(str(semester) if semester else '', s_sub),
        Spacer(1, .4*cm),
    ]

    col_time = 2.1*cm
    col_day  = (usable - col_time) / len(grid['days'])

    header = [Paragraph('Horaire', s_hdr)] + [Paragraph(d['label'], s_hdr) for d in grid['days']]
    data = [header]
    span_cmds = []

    for r_i, row in enumerate(grid['rows'], start=1):
        time_lbl = f"{row['start'].strftime('%Hh%M')}\n{row['end'].strftime('%Hh%M')}"
        line = [Paragraph(time_lbl.replace('\n', '<br/>'), s_time)]
        for c_i, cell in enumerate(row['cells'], start=1):
            if cell is None:
                line.append('')
                continue
            e = cell['entry']
            if e is None:
                line.append('')
            else:
                room_name = e.room.name if e.room else ''
                line2, line3 = {
                    'Classe':     (e.teacher.full_name, room_name),
                    'Enseignant': (e.class_group.name, room_name),
                    'Salle':      (e.class_group.name, e.teacher.full_name),
                }.get(axis_label, (e.teacher.full_name, room_name))
                cell_parts = [f"<b>{_esc(e.subject.title)}</b>", _esc(line2)]
                if line3:
                    cell_parts.append(_esc(line3))
                line.append(Paragraph('<br/>'.join(cell_parts), s_cell))
            if cell['span'] > 1:
                span_cmds.append(('SPAN', (c_i, r_i), (c_i, r_i + cell['span'] - 1)))
        data.append(line)

    tbl = Table(data, colWidths=[col_time] + [col_day] * len(grid['days']), repeatRows=1)
    style_cmds = [
        ('BACKGROUND', (0, 0), (-1, 0), navy),
        ('GRID', (0, 0), (-1, -1), 0.4, border),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('BACKGROUND', (0, 1), (0, -1), lgrey),
    ] + span_cmds
    tbl.setStyle(TableStyle(style_cmds))
    story.append(tbl)

    doc.build(story, onFirstPage=watermark_canvas, onLaterPages=watermark_canvas)
    buf.seek(0)
    return buf


@login_required
def planning_pdf(request):
    """Génère la grille hebdomadaire (PDF, A4 paysage) pour une classe, un enseignant ou une salle."""
    entries, semester, axis_label, axis_obj = _filtered_entries(request)

    if not axis_obj:
        messages.error(request, "Sélectionnez une classe, un enseignant ou une salle pour générer le planning.")
        return redirect('timetable:index')

    grid = build_planning_grid(entries)
    if not grid:
        messages.warning(request, "Aucune séance à afficher pour cette sélection.")
        return redirect('timetable:index')

    from academic_core.pdf_utils import get_institut_config_for_request
    config = get_institut_config_for_request(request)
    buf = generate_planning_pdf_bytes(grid, axis_label, axis_obj, semester, config=config)
    safe = f"EmploiDuTemps_{axis_label}_{axis_obj}".replace(' ', '_').replace('/', '-')
    response = HttpResponse(buf.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{safe}.pdf"'
    return response


@login_required
def planning_excel(request):
    """Génère la grille hebdomadaire (Excel) pour une classe, un enseignant ou une salle."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    entries, semester, axis_label, axis_obj = _filtered_entries(request)

    if not axis_obj:
        messages.error(request, "Sélectionnez une classe, un enseignant ou une salle pour générer le planning.")
        return redirect('timetable:index')

    grid = build_planning_grid(entries)
    if not grid:
        messages.warning(request, "Aucune séance à afficher pour cette sélection.")
        return redirect('timetable:index')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Emploi du temps'

    navy_fill  = PatternFill('solid', fgColor='1E3A5F')
    light_fill = PatternFill('solid', fgColor='F8FAFC')
    thin = Side(style='thin', color='E2E8F0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    if axis_label == 'Classe' and axis_obj is not None:
        filiere_name = axis_obj.program.name if getattr(axis_obj, 'program', None) else ''
        axis_title = f"{filiere_name} - {axis_obj.name}" if filiere_name else axis_obj.name
    else:
        axis_title = f"{axis_label} : {axis_obj}"

    n_cols = len(grid['days']) + 1
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_cell = ws.cell(row=1, column=1, value=f"Emploi du temps — {axis_title}")
    title_cell.font = Font(bold=True, size=13, color='1E3A5F')
    title_cell.alignment = Alignment(horizontal='center')

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    sub_cell = ws.cell(row=2, column=1, value=str(semester) if semester else '')
    sub_cell.font = Font(italic=True, size=9, color='64748B')
    sub_cell.alignment = Alignment(horizontal='center')

    header_row = 4
    ws.cell(row=header_row, column=1, value='Horaire')
    for c_i, d in enumerate(grid['days'], start=2):
        ws.cell(row=header_row, column=c_i, value=d['label'])
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = navy_fill
        cell.alignment = center
        cell.border = border

    for r_i, row in enumerate(grid['rows']):
        xl_row = header_row + 1 + r_i
        time_cell = ws.cell(row=xl_row, column=1,
                             value=f"{row['start'].strftime('%H:%M')}–{row['end'].strftime('%H:%M')}")
        time_cell.font = Font(bold=True, size=9, color='64748B')
        time_cell.alignment = center
        time_cell.fill = light_fill
        time_cell.border = border

        for c_i, cell in enumerate(row['cells'], start=2):
            if cell is None:
                continue
            xl_col = c_i
            e = cell['entry']
            target = ws.cell(row=xl_row, column=xl_col)
            target.border = border
            target.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if e is not None:
                room_name = e.room.name if e.room else ''
                line2, line3 = {
                    'Classe':     (e.teacher.full_name, room_name),
                    'Enseignant': (e.class_group.name, room_name),
                    'Salle':      (e.class_group.name, e.teacher.full_name),
                }.get(axis_label, (e.teacher.full_name, room_name))
                cell_lines = [e.subject.title, line2] + ([line3] if line3 else [])
                target.value = '\n'.join(cell_lines)
                target.font  = Font(size=9, bold=True, color='1E3A5F')
            if cell['span'] > 1:
                ws.merge_cells(start_row=xl_row, start_column=xl_col,
                               end_row=xl_row + cell['span'] - 1, end_column=xl_col)
                for rr in range(xl_row, xl_row + cell['span']):
                    ws.cell(row=rr, column=xl_col).border = border

    ws.column_dimensions['A'].width = 14
    for c_i in range(2, n_cols + 1):
        ws.column_dimensions[get_column_letter(c_i)].width = 22
    for r_i in range(len(grid['rows'])):
        ws.row_dimensions[header_row + 1 + r_i].height = 34

    import io
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe = f"EmploiDuTemps_{axis_label}_{axis_obj}".replace(' ', '_').replace('/', '-')
    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{safe}.xlsx"'
    return response


@login_required
def planning_word(request):
    """Génère la grille hebdomadaire (.doc, ouvert par Word) pour une classe, un enseignant ou une salle."""
    entries, semester, axis_label, axis_obj = _filtered_entries(request)

    if not axis_obj:
        messages.error(request, "Sélectionnez une classe, un enseignant ou une salle pour générer le planning.")
        return redirect('timetable:index')

    grid = build_planning_grid(entries)
    if not grid:
        messages.warning(request, "Aucune séance à afficher pour cette sélection.")
        return redirect('timetable:index')

    from academic_core.pdf_utils import get_institut_config_for_request, get_logo_path
    config    = get_institut_config_for_request(request)
    inst_name = (config.nom if config and getattr(config, 'nom', None) else '') or ''

    import os
    logo_img_tag = ''
    logo_path = get_logo_path(config)
    if logo_path and os.path.exists(logo_path):
        import base64, mimetypes
        mime = mimetypes.guess_type(logo_path)[0] or 'image/png'
        with open(logo_path, 'rb') as _f:
            logo_b64 = base64.b64encode(_f.read()).decode('ascii')
        logo_img_tag = f'<img src="data:{mime};base64,{logo_b64}" height="48" style="margin-bottom:4px;">'

    if axis_label == 'Classe' and axis_obj is not None:
        filiere_name = axis_obj.program.name if getattr(axis_obj, 'program', None) else ''
        axis_title = f"{filiere_name} - {axis_obj.name}" if filiere_name else axis_obj.name
    else:
        axis_title = f"{axis_label.upper()} : {axis_obj}"

    header_cells = ''.join(f'<th>{_esc(d["label"])}</th>' for d in grid['days'])

    body_rows = ''
    for row in grid['rows']:
        time_lbl = f"{row['start'].strftime('%H:%M')}–{row['end'].strftime('%H:%M')}"
        cells_html = ''
        for cell in row['cells']:
            if cell is None:
                continue
            e = cell['entry']
            rowspan = f' rowspan="{cell["span"]}"' if cell['span'] > 1 else ''
            if e is None:
                cells_html += f'<td{rowspan}>&nbsp;</td>'
                continue
            room_name = e.room.name if e.room else ''
            line2, line3 = {
                'Classe':     (e.teacher.full_name, room_name),
                'Enseignant': (e.class_group.name, room_name),
                'Salle':      (e.class_group.name, e.teacher.full_name),
            }.get(axis_label, (e.teacher.full_name, room_name))
            secondary_html = f'<span style="font-size:8pt;color:#475569;">{_esc(line2)}</span>'
            if line3:
                secondary_html += f'<br><span style="font-size:8pt;color:#475569;">{_esc(line3)}</span>'
            cells_html += (
                f'<td{rowspan} style="background:#eff6ff;">'
                f'<b>{_esc(e.subject.title)}</b><br>{secondary_html}</td>'
            )
        body_rows += f'<tr><td style="background:#f8fafc;font-weight:bold;white-space:nowrap;">{time_lbl}</td>{cells_html}</tr>'

    html_content = f"""<html xmlns:o="urn:schemas-microsoft-com:office:office"
      xmlns:w="urn:schemas-microsoft-com:office:word"
      xmlns="http://www.w3.org/TR/REC-html40">
<head><meta charset="UTF-8">
<style>
  body  {{ font-family: Calibri, Arial, sans-serif; font-size: 11pt; margin: 1.5cm; }}
  table {{ border-collapse: collapse; width: 100%; }}
  th, td {{ border: 1px solid #cbd5e1; padding: 5px 8px; font-size: 9pt; text-align: center; vertical-align: middle; }}
  th {{ background: #1e3a5f; color: #fff; }}
  @page {{ mso-footer: f1; }}
</style>
</head>
<body>
<div style="mso-element:footer;" id="f1">
  <p style="text-align:center;font-size:6pt;color:#94A3B8;margin:0;">ERP-RYTAL, App-GEST V1.0</p>
</div>
<div style="text-align:center;margin-bottom:18px;">
  {logo_img_tag}
  <p style="font-size:9pt;color:#64748b;margin:0 0 4px;">{_esc(inst_name.upper())}</p>
  <p style="font-size:14pt;font-weight:bold;color:#1e3a5f;margin:0 0 4px;">EMPLOI DU TEMPS — {_esc(axis_title)}</p>
  <p style="font-size:10pt;color:#64748b;margin:0;">{_esc(semester) if semester else ''}</p>
</div>
<table>
  <tr><th>Horaire</th>{header_cells}</tr>
  {body_rows}
</table>
</body></html>"""

    safe = f"EmploiDuTemps_{axis_label}_{axis_obj}".replace(' ', '_').replace('/', '-')
    response = HttpResponse(html_content, content_type='application/vnd.ms-word')
    response['Content-Disposition'] = f'attachment; filename="{safe}.doc"'
    return response


# ── Envoi par email ──────────────────────────────────────────────────────────

@login_required
def send_class_planning_emails(request):
    """
    Envoie par email, à chaque étudiant inscrit (validé) de la classe
    sélectionnée, l'emploi du temps de sa classe (même PDF pour tous les
    étudiants de la classe — un seul appel à notify_users, contenu partagé).
    """
    if request.method != 'POST' or not request.user.can_manage_dept():
        messages.error(request, "Accès refusé.")
        return redirect('timetable:index')

    from django.urls import reverse
    from academic_core.apps.academic_structure.models import Class as ClassGroup
    from academic_core.apps.students.models import Enrollment
    from academic_core.pdf_utils import get_institut_config_for_request
    from academic_core.apps.notifications.utils import notify_users
    from academic_core.apps.notifications.models import Notification

    class_id = request.POST.get('class_group', '').strip()
    class_group = ClassGroup.objects.filter(pk=class_id).first() if class_id else None
    if not class_group:
        messages.error(request, "Sélectionnez une classe pour envoyer son emploi du temps.")
        return redirect('timetable:index')

    redirect_url = f"{reverse('timetable:index')}?class_group={class_group.pk}"

    entries = TimetableEntry.objects.select_related(
        'class_group', 'subject', 'teacher__user', 'room', 'semester',
    ).filter(is_active=True, class_group=class_group).order_by('day_of_week', 'start_time')
    grid = build_planning_grid(entries)
    if not grid:
        messages.warning(request, f"Aucune séance planifiée pour {class_group.name} — rien à envoyer.")
        return redirect(redirect_url)

    recipients = [
        enr.student.user for enr in
        Enrollment.objects.filter(class_group=class_group, status=Enrollment.STATUS_VALIDATED)
        .select_related('student__user')
        if enr.student.user and enr.student.user.email
    ]
    if not recipients:
        messages.warning(request, f"Aucun étudiant avec adresse email dans {class_group.name}.")
        return redirect(redirect_url)

    config = get_institut_config_for_request(request)
    buf = generate_planning_pdf_bytes(grid, 'Classe', class_group, None, config=config)
    filename = f"EmploiDuTemps_{class_group.name}".replace(' ', '_').replace('/', '-') + '.pdf'

    notify_users(
        recipients=recipients,
        notification_type=Notification.TYPE_TIMETABLE_CHANGE,
        title=f"Emploi du temps — {class_group.name}",
        message=(
            f"L'emploi du temps de votre classe {class_group.name} est disponible, "
            f"en pièce jointe de cet email."
        ),
        send_email=True,
        email_heading=f"Emploi du temps — {class_group.name}",
        email_paragraphs=[
            f"Voici l'emploi du temps actualisé de votre classe {class_group.name}.",
            "Vous le trouverez en pièce jointe de cet email (format PDF).",
        ],
        email_attachments=[(filename, buf.read(), 'application/pdf')],
    )

    messages.success(
        request,
        f"Emploi du temps envoyé par email à {len(recipients)} étudiant(s) de {class_group.name}."
    )
    return redirect(redirect_url)


@login_required
def send_teacher_planning_emails(request):
    """
    Envoie par email, à chaque enseignant ayant au moins une séance active
    dans le périmètre choisi (un département, ou tous les départements de
    l'institut si aucun n'est précisé), son planning personnel — toutes ses
    classes dans ce même périmètre, un PDF différent par enseignant.
    """
    if request.method != 'POST' or not request.user.can_manage_dept():
        messages.error(request, "Accès refusé.")
        return redirect('timetable:index')

    from academic_core.apps.academic_structure.models import Department
    from academic_core.apps.teachers.models import Teacher
    from academic_core.pdf_utils import get_institut_config_for_request
    from academic_core.apps.notifications.utils import notify_users
    from academic_core.apps.notifications.models import Notification

    dept_id = request.POST.get('department', '').strip()
    department = None
    if dept_id:
        department = Department.objects.filter(pk=dept_id).first()
        if not department:
            messages.error(request, "Département introuvable.")
            return redirect('timetable:index')

    faculty = getattr(request, 'active_faculty', None)

    entries_qs = TimetableEntry.objects.select_related(
        'class_group__program__department', 'subject', 'teacher__user', 'room', 'semester',
    ).filter(is_active=True)
    if department:
        entries_qs = entries_qs.filter(class_group__program__department=department)
    elif faculty:
        entries_qs = entries_qs.filter(class_group__program__department__faculty=faculty)

    teacher_ids = entries_qs.values_list('teacher_id', flat=True).distinct()
    teachers = Teacher.objects.filter(pk__in=teacher_ids).select_related('user')

    config = get_institut_config_for_request(request)
    scope_txt = f" au sein du département {department.name}" if department else ""
    sent_count = 0

    for teacher in teachers:
        if not teacher.user or not teacher.user.email:
            continue
        t_entries = entries_qs.filter(teacher=teacher).order_by('day_of_week', 'start_time')
        grid = build_planning_grid(t_entries)
        if not grid:
            continue
        buf = generate_planning_pdf_bytes(grid, 'Enseignant', teacher, None, config=config)
        filename = f"EmploiDuTemps_{teacher.full_name}".replace(' ', '_').replace('/', '-') + '.pdf'

        notify_users(
            recipients=[teacher.user],
            notification_type=Notification.TYPE_TIMETABLE_CHANGE,
            title="Votre emploi du temps",
            message=f"Votre emploi du temps{scope_txt} est disponible, en pièce jointe de cet email.",
            send_email=True,
            email_heading="Votre emploi du temps",
            email_paragraphs=[
                f"Voici votre emploi du temps personnel{scope_txt}, couvrant toutes vos classes.",
                "Vous le trouverez en pièce jointe de cet email (format PDF).",
            ],
            email_attachments=[(filename, buf.read(), 'application/pdf')],
        )
        sent_count += 1

    if sent_count:
        messages.success(
            request,
            f"Emploi du temps envoyé à {sent_count} enseignant(s)"
            + (f" du département {department.name}." if department else " de l'institut.")
        )
    else:
        messages.warning(
            request,
            "Aucun enseignant avec des séances planifiées (et une adresse email) trouvé dans ce périmètre."
        )
    return redirect('timetable:index')
