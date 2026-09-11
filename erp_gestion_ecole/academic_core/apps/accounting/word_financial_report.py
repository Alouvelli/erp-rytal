"""
Export Word (python-docx) des rapports financiers (rapports_financiers.html
:: _rf_compute) — générique à tous les types de rapport (tableau principal
optionnellement à en-têtes colorés par colonne + sections), même famille
visuelle que grades/word_class_results.py.
"""
import io
import os

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from academic_core.pdf_utils import add_docx_watermark, get_logo_path


def _rgb(h):
    h = (h or '0D2244').lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_cell_shd(cell, hex_fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), (hex_fill or '0D2244').lstrip('#'))
    tcPr.append(s)


def _wc(cell, text, bold=False, size=8, color='000000', align='center', fill=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == 'center' else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text) if text is not None else '—')
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = _rgb(color)
    if fill:
        _set_cell_shd(cell, fill)


def _add_table(doc, headers, rows, header_colors=None, default_color='0D2244'):
    if not headers:
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        fill = (header_colors[i] if header_colors and header_colors[i] else default_color)
        _wc(table.rows[0].cells[i], h, bold=True, size=8, color='FFFFFF', fill=fill)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            _wc(table.rows[r].cells[c], val, size=8, fill='FFFFFF' if r % 2 else 'F8FAFC')


def generate_financial_report_word(data, *, label, year, config=None, include_main=True, sections=None):
    """
    `data` : dict retourné par accounting.views._rf_compute(). `include_main`/
    `sections` permettent un export par rubrique (voir _rf_filtered_data) ;
    par défaut (les deux fournis complets) l'export est global.
    """
    sections = sections if sections is not None else data.get('sections', [])

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(42)
    sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(0.8)
    sec.top_margin = sec.bottom_margin = Cm(1.0)

    # ── En-tête : logo institut + nom ─────────────────────────────────────────
    inst_nom = (getattr(config, 'nom', None)) or "Institut Supérieur d'Informatique — ISI"
    logo_path = get_logo_path(config)
    hdr = doc.add_table(rows=1, cols=2)
    hdr.autofit = False
    hdr.columns[0].width = Cm(2.2)
    cell_logo = hdr.cell(0, 0)
    cell_logo.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if logo_path and os.path.exists(logo_path):
        p_logo.add_run().add_picture(logo_path, height=Cm(1.4))
    cell_nom = hdr.cell(0, 1)
    cell_nom.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p_nom = cell_nom.paragraphs[0]
    run_nom = p_nom.add_run(inst_nom)
    run_nom.bold = True
    run_nom.font.size = Pt(11)
    run_nom.font.color.rgb = _rgb('0D2244')
    doc.add_paragraph()

    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tp.add_run(label.upper())
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = _rgb('0D2244')

    sp = doc.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = sp.add_run(f"Année académique : {year.label if year else '—'}")
    run2.font.size = Pt(9)
    run2.font.color.rgb = _rgb('64748B')

    if data.get('summary'):
        kp = doc.add_paragraph()
        kp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run3 = kp.add_run('   |   '.join(f"{v} {lbl}" for lbl, v in data['summary']))
        run3.font.size = Pt(9)
        run3.font.color.rgb = _rgb('334155')

    if include_main:
        header_cells = data.get('header_cells')
        if header_cells:
            headers = [c['label'] for c in header_cells]
            header_colors = [c.get('color') for c in header_cells]
        else:
            headers = data.get('columns', [])
            header_colors = None
        _add_table(doc, headers, data.get('rows', []), header_colors)

    for section in sections:
        doc.add_paragraph()
        hp = doc.add_paragraph()
        run4 = hp.add_run(section.get('title', ''))
        run4.bold = True
        run4.font.size = Pt(11)
        run4.font.color.rgb = _rgb('0D2244')
        if section.get('summary'):
            kp2 = doc.add_paragraph()
            run5 = kp2.add_run('   |   '.join(f"{v} {lbl}" for lbl, v in section['summary']))
            run5.font.size = Pt(9)
            run5.font.color.rgb = _rgb('334155')
        _add_table(doc, section.get('columns', []), section.get('rows', []),
                   default_color=(section.get('color') or '0D2244').lstrip('#'))

    add_docx_watermark(doc)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
