"""Génération PDF/Excel du module Rapports COIP — les deux formats
partagent EXACTEMENT les mêmes indicateurs (services.compute_coip_report_kpis),
contrairement à GestionCOIP où l'Excel n'affichait que 3 indicateurs contre 6
pour le PDF (incohérence relevée à l'audit et corrigée ici).

Pour un rapport de type ANNUEL, les deux formats ajoutent en plus le détail
complet de TOUTES les activités de la COIP sur la période
(services.compute_coip_annual_activities) et le taux d'insertion par filière
(services.compute_insertion_by_filiere) — un rapport mensuel/trimestriel reste
volontairement synthétique (juste les KPI), l'annuel est le document exhaustif
destiné à la direction/tutelle."""
import io

from .models import Report
from .services import (
    compute_coip_annual_activities, compute_coip_report_kpis,
    compute_insertion_by_filiere,
)

# ── Sections détaillées (rapport ANNUEL uniquement) ─────────────────────────
# (clé dans compute_coip_annual_activities, titre de section, colonnes, extracteur de ligne)
_ACTIVITY_SECTIONS = [
    (
        'activities', 'Activités COIP organisées',
        ['Titre', 'Type', 'Date', 'Lieu', 'Statut', 'Participants'],
        lambda a: [a.titre, a.get_type_activite_display(), a.date_debut.strftime('%d/%m/%Y'),
                   a.lieu or '—', a.get_status_display(), a.participants.count()],
    ),
    (
        'partnerships', 'Conventions de partenariat signées',
        ['Intitulé', 'Partenaire', 'Signature', 'Expiration', 'Statut'],
        lambda p: [p.intitule, p.partner.raison_sociale, p.date_signature.strftime('%d/%m/%Y'),
                   p.date_expiration.strftime('%d/%m/%Y') if p.date_expiration else '—', p.get_status_display()],
    ),
    (
        'internships', 'Stages débutés',
        ['Étudiant', 'Intitulé', 'Partenaire', 'Début', 'Statut'],
        lambda i: [str(i.student), i.titre, i.partner.raison_sociale, i.date_debut.strftime('%d/%m/%Y'), i.get_status_display()],
    ),
    (
        'visits', 'Sorties pédagogiques',
        ['Intitulé', 'Lieu', 'Départ', 'Responsable', 'Statut'],
        lambda v: [v.intitule, v.lieu, v.date_depart.strftime('%d/%m/%Y'),
                   v.responsable.get_full_name() if v.responsable else '—', v.get_status_display()],
    ),
    (
        'orientation_sessions', "Séances d'orientation",
        ['Étudiant', 'Type', 'Date', 'Conseiller'],
        lambda s: [str(s.student), s.get_type_session_display(), s.date_session.strftime('%d/%m/%Y'),
                   s.conseiller.get_full_name() if s.conseiller else '—'],
    ),
    (
        'recommendations', 'Demandes de recommandation',
        ['Étudiant', 'Destinataire', 'Date de demande', 'Statut'],
        lambda r: [str(r.student), r.destinataire or '—', r.date_demande.strftime('%d/%m/%Y'), r.get_status_display()],
    ),
    (
        'opportunities', 'Opportunités publiées',
        ['Titre', 'Type', 'Filière ciblée', 'Statut'],
        lambda o: [o.titre, o.get_type_opportunite_display(), o.filiere_cible.name if o.filiere_cible else 'Toutes', o.get_status_display()],
    ),
]

_KPI_LABELS = [
    ('etudiants_suivis', 'Étudiants suivis (orientation, recommandation ou stage)'),
    ('nouveaux_alumni', 'Nouveaux alumni enregistrés'),
    ('partenariats_actifs', 'Partenariats actifs'),
    ('stages_realises', 'Stages achevés sur la période'),
    ('activites_organisees', 'Activités organisées'),
    ('seances_orientation', "Séances d'orientation tenues"),
    ('recommandations_livrees', 'Recommandations livrées'),
]


def generate_coip_report_pdf(request, report):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    from academic_core.pdf_utils import (
        get_institut_config_for_request, logo_image, senegal_flag, watermark_canvas,
    )

    kpis = compute_coip_report_kpis(report.periode_debut, report.periode_fin)
    inst_cfg = get_institut_config_for_request(request)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1.7*cm, leftMargin=1.7*cm, topMargin=1.7*cm, bottomMargin=1.7*cm)

    navy = colors.HexColor('#00173B')
    gold = colors.HexColor('#D4AF37')
    light = colors.HexColor('#EFF6FF')
    ss = getSampleStyleSheet()
    sc = ParagraphStyle
    s_ttl = sc('rp_ttl', parent=ss['Normal'], alignment=TA_CENTER, fontSize=16, fontName='Helvetica-Bold', textColor=navy)
    s_sub = sc('rp_sub', parent=ss['Normal'], alignment=TA_CENTER, fontSize=10, textColor=colors.HexColor('#64748B'))
    s_kpi_label = sc('rp_kl', parent=ss['Normal'], fontSize=10.5, fontName='Helvetica-Bold', textColor=navy)
    s_kpi_val = sc('rp_kv', parent=ss['Normal'], fontSize=14, fontName='Helvetica-Bold', alignment=TA_CENTER)

    inst_nom = inst_cfg.nom if inst_cfg else ''
    inst_sigle = (inst_cfg.sigle if inst_cfg else '') or inst_nom
    flag_el = senegal_flag(width=2.2*cm, height=1.5*cm)
    logo_el = logo_image(width=2.2*cm, height=1.5*cm, config=inst_cfg) or Paragraph(
        inst_sigle, sc('rp_lb', parent=ss['Normal'], alignment=TA_CENTER, fontSize=9, fontName='Helvetica-Bold', textColor=navy))
    rep_txt = Paragraph(
        '<b>REPUBLIQUE DU SÉNÉGAL</b><br/><font size="8">Un Peuple — Un But — Une Foi</font>'
        f'<br/><br/><font size="13"><b>{inst_sigle}</b></font><br/><font size="9">{inst_nom}</font>',
        sc('rp_h', parent=ss['Normal'], alignment=TA_CENTER, leading=14),
    )
    hdr_tbl = Table([[flag_el, rep_txt, logo_el]], colWidths=[3*cm, 12.6*cm, 3*cm])
    hdr_tbl.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))

    story = [
        hdr_tbl, Spacer(1, .3*cm),
        HRFlowable(width='100%', thickness=2, color=navy),
        HRFlowable(width='100%', thickness=1, color=gold, spaceAfter=14),
        Spacer(1, .3*cm),
        Paragraph("Cellule d'Orientation et d'Insertion Professionnelle", s_sub),
        Paragraph(report.titre, s_ttl),
        Paragraph(
            f"Période : {report.periode_debut.strftime('%d/%m/%Y')} — {report.periode_fin.strftime('%d/%m/%Y')} · "
            f"{report.get_type_rapport_display()}",
            s_sub,
        ),
        Spacer(1, .8*cm),
    ]

    rows = []
    for i in range(0, len(_KPI_LABELS), 2):
        pair = _KPI_LABELS[i:i+2]
        row = []
        for key, label in pair:
            cell = Table(
                [[Paragraph(str(kpis[key]), s_kpi_val)], [Paragraph(label, sc('rp_kl2', parent=s_kpi_label, alignment=TA_CENTER, fontSize=9))]],
                colWidths=[8.5*cm],
            )
            cell.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), light),
                ('BOX', (0, 0), (-1, -1), 0.5, navy),
                ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ]))
            row.append(cell)
        if len(row) == 1:
            row.append('')
        rows.append(row)

    kpi_table = Table(rows, colWidths=[8.7*cm, 8.7*cm], spaceBefore=4, spaceAfter=4)
    kpi_table.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BOTTOMPADDING', (0, 0), (-1, -1), 10)]))
    story.append(kpi_table)

    if report.type_rapport == Report.TYPE_ANNUEL:
        s_section = sc('rp_section', parent=ss['Normal'], fontSize=12, fontName='Helvetica-Bold', textColor=navy, spaceBefore=16, spaceAfter=6)
        s_cell = sc('rp_cell', parent=ss['Normal'], fontSize=8, leading=10)
        s_head = sc('rp_head', parent=ss['Normal'], fontSize=8.5, fontName='Helvetica-Bold', textColor=colors.white)
        s_empty = sc('rp_empty', parent=ss['Normal'], fontSize=9, textColor=colors.HexColor('#64748B'), alignment=TA_LEFT)

        story.append(Spacer(1, .4*cm))
        story.append(HRFlowable(width='100%', thickness=1, color=gold))
        story.append(Paragraph('Détail complet des activités de la COIP sur la période', s_section))

        activities_data = compute_coip_annual_activities(report.periode_debut, report.periode_fin)
        for key, section_title, headers, extractor in _ACTIVITY_SECTIONS:
            items = activities_data[key]
            story.append(Paragraph(f"{section_title} ({len(items)})", s_section))
            if not items:
                story.append(Paragraph('Aucune donnée sur cette période.', s_empty))
                continue
            table_data = [[Paragraph(h, s_head) for h in headers]]
            for item in items:
                table_data.append([Paragraph(str(v), s_cell) for v in extractor(item)])
            tbl = Table(table_data, repeatRows=1)
            tbl.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), navy),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, light]),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]))
            story.append(tbl)

        # Insertion par filière (photo globale des alumni au moment du rapport
        # — pas restreinte à la période, contrairement aux sections ci-dessus,
        # car l'insertion se mesure sur la cohorte des diplômés, pas sur les
        # évènements COIP de l'année).
        insertion = compute_insertion_by_filiere()
        story.append(Paragraph(
            f"Taux d'insertion par filière (toutes promotions — {insertion['taux_global']} % global)", s_section,
        ))
        if insertion['par_filiere']:
            ins_data = [[Paragraph(h, s_head) for h in ['Filière', 'Alumni recensés', 'En activité', 'Taux']]]
            for r in insertion['par_filiere']:
                ins_data.append([Paragraph(str(v), s_cell) for v in [r['filiere_nom'], r['total'], r['employed'], f"{r['taux']} %"]])
            ins_tbl = Table(ins_data, repeatRows=1)
            ins_tbl.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), navy),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, light]),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]))
            story.append(ins_tbl)
        else:
            story.append(Paragraph('Aucun alumni recensé.', s_empty))

    doc.build(story, onFirstPage=watermark_canvas, onLaterPages=watermark_canvas)
    buffer.seek(0)
    return buffer


def generate_coip_report_excel(report):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    kpis = compute_coip_report_kpis(report.periode_debut, report.periode_fin)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Rapport COIP'

    navy = '00173B'
    thin = Side(style='thin', color='CBD5E1')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells('A1:B1')
    ws['A1'] = report.titre
    ws['A1'].font = Font(bold=True, color='FFFFFF', size=13)
    ws['A1'].fill = PatternFill('solid', fgColor=navy)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26

    ws.merge_cells('A2:B2')
    ws['A2'] = f"Période : {report.periode_debut.strftime('%d/%m/%Y')} — {report.periode_fin.strftime('%d/%m/%Y')}"
    ws['A2'].font = Font(italic=True, color='64748B', size=9)
    ws['A2'].alignment = Alignment(horizontal='center')

    headers = ['Indicateur', 'Valeur']
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color='FFFFFF', size=10)
        c.fill = PatternFill('solid', fgColor='1E3A5F')
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border
    ws.column_dimensions['A'].width = 55
    ws.column_dimensions['B'].width = 14

    rownum = 5
    for key, label in _KPI_LABELS:
        ws.cell(row=rownum, column=1, value=label).border = border
        c = ws.cell(row=rownum, column=2, value=kpis[key])
        c.border = border
        c.alignment = Alignment(horizontal='center')
        rownum += 1

    if report.type_rapport == Report.TYPE_ANNUEL:
        def _write_table_sheet(title, headers, rows):
            sheet = wb.create_sheet(title=title[:31])  # 31 car. max (limite Excel)
            for i, h in enumerate(headers, 1):
                c = sheet.cell(row=1, column=i, value=h)
                c.font = Font(bold=True, color='FFFFFF', size=10)
                c.fill = PatternFill('solid', fgColor='1E3A5F')
                c.alignment = Alignment(horizontal='center', vertical='center')
                c.border = border
            for r_idx, row in enumerate(rows, 2):
                for c_idx, value in enumerate(row, 1):
                    cell = sheet.cell(row=r_idx, column=c_idx, value=value)
                    cell.border = border
            for i in range(1, len(headers) + 1):
                sheet.column_dimensions[get_column_letter(i)].width = 24
            if not rows:
                sheet.cell(row=2, column=1, value='Aucune donnée sur cette période.').font = Font(italic=True, color='64748B')

        activities_data = compute_coip_annual_activities(report.periode_debut, report.periode_fin)
        for key, section_title, headers, extractor in _ACTIVITY_SECTIONS:
            rows = [extractor(item) for item in activities_data[key]]
            _write_table_sheet(section_title, headers, rows)

        insertion = compute_insertion_by_filiere()
        insertion_rows = [
            [r['filiere_nom'], r['total'], r['employed'], f"{r['taux']} %"]
            for r in insertion['par_filiere']
        ]
        _write_table_sheet(
            "Insertion par filière", ['Filière', 'Alumni recensés', 'En activité', 'Taux'], insertion_rows,
        )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
