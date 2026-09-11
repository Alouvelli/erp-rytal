"""
Utilitaires partagés pour la génération de PDFs.
Fournit le logo de l'institut et le drapeau du Sénégal en éléments ReportLab.
"""
import math
import os

from reportlab.graphics.shapes import Drawing, Rect, Polygon
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Image, Paragraph

# Couleur/style partagés du filigrane "ERP-RYTAL" apposé en pied de TOUS les
# fichiers PDF/Word téléchargés ou imprimés générés par la plateforme.
WATERMARK_TEXT  = 'ERP-RYTAL, App-GEST V1.0'
WATERMARK_COLOR = colors.HexColor('#94A3B8')

# Chemin absolu vers le logo ISI (fallback statique)
_BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LOGO_PATH = os.path.join(_BASE_DIR, 'static', 'img', 'logo_isi.png')

# Couleurs du drapeau sénégalais
_FLAG_GREEN  = colors.HexColor('#00853F')
_FLAG_YELLOW = colors.HexColor('#FDEF42')
_FLAG_RED    = colors.HexColor('#E31B23')
_STAR_GREEN  = colors.HexColor('#00853F')


def get_logo_path(config=None):
    """
    Retourne le chemin absolu du logo à utiliser :
    - Si config (InstitutConfig) est fourni et a un logo → chemin media
    - Sinon → chemin du logo statique ISI (fallback)
    """
    if config is not None:
        logo_field = getattr(config, 'logo', None)
        if logo_field and logo_field.name:
            try:
                path = logo_field.path
                if os.path.exists(path):
                    return path
            except Exception:
                pass
    return LOGO_PATH if os.path.exists(LOGO_PATH) else None


def logo_image(width=2.2 * cm, height=1.4 * cm, config=None):
    """
    Retourne un élément Image ReportLab du logo de l'institut.
    Utilise le logo de l'InstitutConfig si fourni, sinon le logo statique ISI.
    """
    path = get_logo_path(config)
    if path:
        return Image(path, width=width, height=height, kind='proportional')
    return None


def get_institut_config_for_request(request):
    """Helper : retourne l'InstitutConfig de l'institut actif depuis une requête."""
    faculty = getattr(request, 'active_faculty', None)
    if not faculty:
        return None
    try:
        from academic_core.apps.academic_structure.models import InstitutConfig
        return InstitutConfig.objects.using('default').get(faculty=faculty)
    except Exception:
        return None


def watermark_paragraph():
    """
    DÉCONSEILLÉ pour un filigrane de PIED DE PAGE : un Paragraph ajouté au
    `story` ne s'affiche qu'à la suite du dernier contenu, pas au bas physique
    de la page (sur un document court, il apparaît donc au milieu de la page,
    pas en bas). Utiliser `watermark_canvas` en callback
    `doc.build(story, onFirstPage=watermark_canvas, onLaterPages=watermark_canvas)`
    à la place, qui ancre le filigrane au bas de CHAQUE page quelle que soit la
    longueur du contenu. Conservé uniquement pour compatibilité.
    """
    style = ParagraphStyle(
        'erp_rytal_watermark', alignment=TA_CENTER,
        fontSize=6, textColor=WATERMARK_COLOR,
    )
    return Paragraph(WATERMARK_TEXT, style)


def draw_watermark(canvas_obj, page_width, y=0.4 * cm, font_size=6):
    """
    Dessine directement le filigrane "ERP-RYTAL" centré en bas de page sur un
    canvas ReportLab bas niveau (`reportlab.pdfgen.canvas.Canvas`), à appeler
    avant `canvas.showPage()`/`canvas.save()`. Utile pour les documents qui ne
    passent pas par Platypus/`SimpleDocTemplate` (ex. carte étudiant recto/verso).
    """
    canvas_obj.saveState()
    canvas_obj.setFont('Helvetica', font_size)
    canvas_obj.setFillColor(WATERMARK_COLOR)
    canvas_obj.drawCentredString(page_width / 2, y, WATERMARK_TEXT)
    canvas_obj.restoreState()


def watermark_canvas(canvas_obj, doc):
    """
    Callback `onFirstPage`/`onLaterPages` pour un document Platypus
    (`doc.build(story, onFirstPage=watermark_canvas, onLaterPages=watermark_canvas)`)
    — dessine le filigrane "ERP-RYTAL" en bas de CHAQUE page (contrairement à
    `watermark_paragraph()`, qui n'apparaît qu'en fin de la dernière page).
    Ne pas combiner les deux sur un même document (redondant).
    """
    draw_watermark(canvas_obj, doc.pagesize[0])


def add_docx_watermark(doc):
    """
    Appose le filigrane "ERP-RYTAL" dans le PIED DE PAGE réel (`section.footer`)
    d'un document python-docx (`docx.Document`), à appeler à tout moment avant
    `doc.save(...)`. Contrairement à un simple `doc.add_paragraph(...)` (qui
    s'insère dans le flux du corps du texte et n'atterrit en bas de page que si
    le contenu remplit la page jusque-là), un pied de page Word reste ancré au
    bas de CHAQUE page, quelle que soit la longueur du contenu.
    """
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for section in doc.sections:
        section.footer.is_linked_to_previous = False
        para = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
        para.text = ''
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(WATERMARK_TEXT)
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)


def simple_list_pdf(request, *, title, headers, rows, filename,
                     subtitle=None, col_widths=None, landscape_mode=True):
    """
    Génère un PDF tabulaire simple (logo + titre + tableau + filigrane
    ERP-RYTAL) pour l'export « Télécharger PDF » d'une liste — utilisé par le
    Plan Stratégique (Cadrages/Axes/Objectifs/Actions/Activités/Indicateurs)
    et réutilisable pour toute autre liste simple. `rows` : liste de listes
    de valeurs déjà formatées en texte (une ligne par enregistrement).
    """
    import io
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.pagesizes import landscape as _landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )
    from django.http import HttpResponse
    from django.utils.html import escape

    config = get_institut_config_for_request(request)
    pagesize = _landscape(A4) if landscape_mode else A4
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=pagesize,
        leftMargin=1.3 * cm, rightMargin=1.3 * cm, topMargin=1.2 * cm, bottomMargin=1.4 * cm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'erp_list_title', parent=styles['Heading1'], alignment=TA_CENTER,
        fontSize=15, textColor=colors.HexColor('#003D82'),
    )
    subtitle_style = ParagraphStyle(
        'erp_list_subtitle', parent=styles['Normal'], alignment=TA_CENTER,
        fontSize=9, textColor=colors.HexColor('#64748B'),
    )
    cell_style = ParagraphStyle('erp_list_cell', parent=styles['Normal'], fontSize=8, leading=10)
    header_style = ParagraphStyle(
        'erp_list_header', parent=styles['Normal'], fontSize=8.5, leading=10,
        textColor=colors.white, fontName='Helvetica-Bold',
    )

    institut_nom = (config.nom or config.sigle) if config else 'Plateforme de Gestion Académique'
    logo = logo_image(width=1.8 * cm, height=1.8 * cm, config=config)
    story = []
    if logo:
        head_tbl = Table([[logo, Paragraph(f"<b>{escape(institut_nom)}</b>", title_style)]],
                          colWidths=[2.4 * cm, None])
        head_tbl.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN', (1, 0), (1, 0), 'CENTER'),
        ]))
        story.append(head_tbl)
    else:
        story.append(Paragraph(escape(institut_nom), title_style))
    story.append(Paragraph(escape(title), title_style))
    if subtitle:
        story.append(Paragraph(escape(subtitle), subtitle_style))
    story.append(Spacer(1, 0.5 * cm))

    table_data = [[Paragraph(escape(str(h)), header_style) for h in headers]]
    for row in rows:
        table_data.append([Paragraph(escape('' if v is None else str(v)), cell_style) for v in row])

    tbl = Table(table_data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#003D82')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(tbl)

    doc.build(story, onFirstPage=watermark_canvas, onLaterPages=watermark_canvas)
    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def senegal_flag(width=2.4 * cm, height=1.6 * cm):
    """
    Retourne un Drawing ReportLab représentant le drapeau du Sénégal :
    trois bandes verticales vert / jaune / rouge avec une étoile verte au centre.
    """
    d = Drawing(width, height)
    stripe = width / 3.0

    # Trois bandes verticales
    d.add(Rect(0,          0, stripe, height, fillColor=_FLAG_GREEN,  strokeColor=None))
    d.add(Rect(stripe,     0, stripe, height, fillColor=_FLAG_YELLOW, strokeColor=None))
    d.add(Rect(stripe * 2, 0, stripe, height, fillColor=_FLAG_RED,    strokeColor=None))

    # Étoile verte à 5 branches centrée dans la bande jaune
    cx = width / 2.0
    cy = height / 2.0
    r_outer = height * 0.28
    r_inner = r_outer * 0.42

    pts = []
    for i in range(10):
        r = r_outer if i % 2 == 0 else r_inner
        # Commence au sommet (90°) et tourne dans le sens anti-horaire
        angle = math.pi / 2 + i * math.pi / 5
        pts.append(cx + r * math.cos(angle))
        pts.append(cy + r * math.sin(angle))

    d.add(Polygon(pts, fillColor=_STAR_GREEN, strokeColor=None))

    return d
